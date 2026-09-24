"""Dashboard del administrador del sistema de asistencia.

Ventana principal (1120x700) con los módulos Inicio, Asistencia, Reportes y
Configuración. Se conecta a Firestore, consulta las marcaciones del día y
actualiza los indicadores periódicamente (auto-refresco cada POLL_INTERVALO).
Incluye los widgets de navegación lateral (NavButton) y de calendario
(CalendarWidget) reutilizables por el resto del sistema.
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import sys
import os
import threading
from datetime import datetime, timedelta

from modelos import get_db, Alerta, FERIADOS
from google.cloud.firestore_v1.base_query import FieldFilter

from theme import COLORS as C, FONTS as F, configure_styles
from components import (
    SurfaceCard, PageHeader, MetricCard, StatusBadge, PrimaryButton,
    SecondaryButton, EmptyState, Toast,
)
import icons


def ruta_relativa(ruta):
    """Devuelve la ruta absoluta de un recurso dentro del proyecto.

    Compatible con ejecutables generados por PyInstaller: en modo "frozen"
    resuelve desde el directorio temporal (_MEIPASS).
    """
    if getattr(sys, "frozen", False):
        base = sys._MEIPASS
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, ruta)


def configurar_icono(ventana):
    """Asigna el ícono de la aplicación a la ventana según la plataforma."""
    try:
        if sys.platform == "win32":
            ventana.iconbitmap(ruta_relativa("icon/Fixmol_icon.ico"))
        else:
            from PIL import Image, ImageTk
            icono = Image.open(ruta_relativa("img/Fixmol3.png"))
            icono = icono.resize((64, 64), Image.LANCZOS)
            foto = ImageTk.PhotoImage(icono)
            ventana.iconphoto(True, foto)
            ventana._icono_app = foto
    except Exception:
        pass


ANCHO = 1120
ALTO = 700
ALTO_TOPBAR = 38
ANCHO_SIDEBAR = 244

HORA_ENTRADA_DEFECTO = "09:30"
HORA_SALIDA_DEFECTO = "17:30"
POLL_INTERVALO = 60000


def _minutos_diff(hora_menor, hora_mayor):
    """Diferencia en minutos entre dos horas HH:MM o HH:MM:SS (nunca negativa).

    Args:
        hora_menor (str): Hora inicial en formato HH:MM o HH:MM:SS.
        hora_mayor (str): Hora final en formato HH:MM o HH:MM:SS.

    Returns:
        int: Minutos de diferencia; 0 si el formato no es válido.
    """

    def _a_minutos(valor):
        """Convierte HH:MM[ :SS] a minutos totales; None si es inválido."""
        try:
            partes = valor.split(":")
            return int(partes[0]) * 60 + int(partes[1])
        except (ValueError, IndexError):
            return None

    t1 = _a_minutos(hora_menor)
    t2 = _a_minutos(hora_mayor)
    if t1 is None or t2 is None:
        return 0
    return max(0, t2 - t1)


def _es_dia_habl(fecha_str):
    """Indica si la fecha dada (YYYY-MM-DD) es día hábil y no feriado."""
    if fecha_str in FERIADOS:
        return False
    try:
        d = datetime.strptime(fecha_str, "%Y-%m-%d")
    except Exception:
        return False
    return d.weekday() < 5


def _hora_valida(valor):
    """Valida que un texto sea una hora en formato HH:MM."""
    try:
        datetime.strptime(valor, "%H:%M")
        return True
    except ValueError:
        return False


class NavButton(tk.Frame):
    """Botón de navegación del sidebar con ícono, texto y estado activo/hover."""

    def __init__(self, master, texto, icono, command, tamano=(216, 56)):
        """Inicializa el botón del sidebar con ícono, texto y comando de clic."""
        super().__init__(master, bg=C["bg_sidebar"], height=tamano[1],
                         cursor="hand2")
        self.pack_propagate(False)
        self.texto = texto
        self.icono = icono
        self.command = command
        self.activo = False
        self.hover = False

        self._ind = None
        self._icon = icons.crear_icono(icono, size=22, color=C["text_secondary"],
                                       bg=C["bg_sidebar"], master=self)
        self._icon.place(x=12, y=17)

        self._lbl = tk.Label(self, text=texto, bg=C["bg_sidebar"],
                             fg=C["text_secondary"], font=F["nav"])
        self._lbl.place(x=52, y=19, anchor="w")

        self._vincular(self)
        self.bind("<Button-1>", self._click)

    def _vincular(self, w):
        """Vincula clic y hover a un widget y a todos sus hijos."""
        for hijo in w.winfo_children():
            self._vincular(hijo)
        w.bind("<Button-1>", self._click)
        w.bind("<Enter>", lambda e: self.set_hover(True))
        w.bind("<Leave>", lambda e: self.set_hover(False))

    def _click(self, e):
        """Ejecuta el comando asociado al botón."""
        self.command()

    def _redibujar(self):
        """Redibuja los colores, tipografía e ícono según el estado del botón."""
        color_fg = C["text_primary"] if (self.activo or self.hover) else C["text_secondary"]
        bg = C["surface_raised"] if self.activo else (C["surface_hover"] if self.hover else C["bg_sidebar"])
        fuente = F["nav_active"] if (self.activo or self.hover) else F["nav"]
        self.config(bg=bg)
        self._lbl.config(bg=bg, fg=color_fg, font=fuente)
        self._icon.destroy()
        self._icon = icons.crear_icono(
            self.icono, size=22, color=color_fg, bg=bg, master=self)
        self._icon.place(x=12, y=17)
        if self.activo and not self._ind:
            self._ind = tk.Frame(self, bg=C["accent"], width=4)
            self._ind.place(x=0, y=0, relheight=1)
        elif not self.activo and self._ind:
            self._ind.destroy()
            self._ind = None

    def set_activo(self, valor):
        """Define el estado activo del botón y lo redibuja."""
        self.activo = bool(valor)
        self._redibujar()

    def set_hover(self, valor):
        """Define el estado de hover del botón y lo redibuja."""
        self.hover = bool(valor)
        self._redibujar()


class CalendarWidget(tk.Frame):
    """Calendario mensual con selección de día y marcas de color por fecha."""

    NOMBRES_DIAS = ["Lu", "Ma", "Mi", "Ju", "Vi", "S\u00e1", "Do"]
    NOMBRES_MES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                   "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre",
                   "Diciembre"]

    def __init__(self, master, seleccionada=None, on_seleccion=None, marcas=None):
        """Inicializa el calendario con la fecha seleccionada, callback y marcas."""
        super().__init__(master, bg=C["surface"], highlightthickness=1,
                         highlightbackground=C["border"], width=330, height=520)
        self.pack_propagate(False)
        self.seleccionada = seleccionada or datetime.now().date()
        self.mostrando = self.seleccionada
        self.on_seleccion = on_seleccion
        self.marcas = marcas or {}
        self._celdas = {}
        self._dibujar()

    def _dibujar(self):
        """Construye la navegación, la cabecera de días, la cuadrícula y la leyenda."""
        for w in self.winfo_children():
            w.destroy()
        self._celdas = {}

        nav = tk.Frame(self, bg=C["surface"])
        nav.pack(fill="x", padx=10, pady=(12, 6))

        tk.Button(nav, text="\u2039", bg=C["surface_raised"], fg=C["text_primary"],
                  activebackground=C["surface_hover"], activeforeground=C["text_primary"],
                  relief="flat", font=F["body_bold"], width=3, cursor="hand2",
                  command=lambda: self._mover_mes(-1)).pack(side="left")

        self._lbl_mes = tk.Label(nav, text="", bg=C["surface"], fg=C["text_primary"],
                                 font=F["body_bold"])
        self._lbl_mes.pack(side="left", expand=True)

        hoy_btn = tk.Button(nav, text="Hoy", bg=C["accent"], fg=C["white"],
                            activebackground=C["accent_hover"],
                            activeforeground=C["white"], relief="flat",
                            font=F["small_bold"], padx=10, cursor="hand2",
                            command=lambda: self._ir_hoy())
        hoy_btn.pack(side="left", padx=(4, 0))

        tk.Button(nav, text="\u203a", bg=C["surface_raised"], fg=C["text_primary"],
                  activebackground=C["surface_hover"], activeforeground=C["text_primary"],
                  relief="flat", font=F["body_bold"], width=3, cursor="hand2",
                  command=lambda: self._mover_mes(1)).pack(side="left", padx=(4, 0))

        cabecera = tk.Frame(self, bg=C["surface"])
        cabecera.pack(fill="x", padx=10)
        for i, dia in enumerate(self.NOMBRES_DIAS):
            tk.Label(cabecera, text=dia, bg=C["surface"], fg=C["text_muted"],
                     font=F["small_bold"], width=5).grid(
                row=0, column=i, padx=2, pady=(0, 6))

        self._cuadricula = tk.Frame(self, bg=C["surface"])
        self._cuadricula.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        self._llenar_cuadricula()

        leyenda = tk.Frame(self, bg=C["surface"])
        leyenda.pack(fill="x", padx=12, pady=(0, 12))
        for txt, color in [("Presente", C["success"]), ("Atraso", C["warning"]),
                           ("Ausente", C["danger"])]:
            tk.Canvas(leyenda, width=7, height=7, bg=C["surface"],
                      highlightthickness=0).pack(side="left", padx=(0, 3))
            c = tk.Canvas(leyenda, width=7, height=7, bg=C["surface"],
                          highlightthickness=0, bd=0)
            c.create_oval(0, 0, 7, 7, fill=color, outline="")
            c.pack(side="left", padx=(4, 3))
            tk.Label(leyenda, text=txt, bg=C["surface"], fg=C["text_secondary"],
                     font=F["small"]).pack(side="left", padx=(0, 10))

        self._actualizar_titulo()

    def _mover_mes(self, delta):
        """Desplaza el mes visible: -1 para el anterior, 1 para el siguiente."""
        mes = self.mostrando.month - 1 + delta
        anio = self.mostrando.year + (mes // 12)
        mes = mes % 12 + 1
        self.mostrando = self.mostrando.replace(year=anio, month=mes)
        self._llenar_cuadricula()
        self._actualizar_titulo()

    def _ir_hoy(self):
        """Selecciona el día de hoy y notifica al callback de selección."""
        hoy = datetime.now().date()
        self.mostrando = hoy
        self.seleccionada = hoy
        self._llenar_cuadricula()
        self._actualizar_titulo()
        if self.on_seleccion:
            self.on_seleccion(hoy.isoformat())

    def _celda_click(self, dia, anio, mes):
        """Selecciona un día al hacer clic y notifica al callback de selección."""
        try:
            fecha = datetime(anio, mes, dia).date()
        except ValueError:
            return
        self.seleccionada = fecha
        self._llenar_cuadricula()
        if self.on_seleccion:
            self.on_seleccion(fecha.isoformat())

    def _actualizar_titulo(self):
        """Actualiza el rótulo "Mes Año" de la cabecera del calendario."""
        self._lbl_mes.config(
            text=f"{self.NOMBRES_MES[self.mostrando.month - 1]} {self.mostrando.year}")

    def _llenar_cuadricula(self):
        """Rellena las celdas del mes mostrado (selección, hoy y marcas de color)."""
        for w in self._cuadricula.winfo_children():
            w.destroy()
        self._celdas = {}
        for fila in range(6):
            self._cuadricula.rowconfigure(fila, weight=1)
        for col in range(7):
            self._cuadricula.columnconfigure(col, weight=1)
        mes = self.mostrando.month
        anio = self.mostrando.year
        primero = datetime(anio, mes, 1)
        lunes = primero.weekday()
        dias_mes = 31 if mes in (1, 3, 5, 7, 8, 10, 12) else (30 if mes in (4, 6, 9, 11) else (29 if (anio % 4 == 0 and anio % 100 != 0) or anio % 400 == 0 else 28))
        hoy = datetime.now().date()
        for i in range(42):
            fila = i // 7
            col = i % 7
            num_dia = i - lunes + 1
            celda = tk.Frame(self._cuadricula, bg=C["surface"])
            celda.grid(row=fila, column=col, sticky="nsew", padx=1, pady=1)
            if 1 <= num_dia <= dias_mes:
                fecha = datetime(anio, mes, num_dia).date()
                es_hoy = fecha == hoy
                es_sel = fecha == self.seleccionada
                cont = tk.Frame(celda, bg=C["surface"])
                cont.place(relx=0.5, rely=0.44, anchor="center")
                cir = tk.Canvas(cont, width=40, height=40, bg=C["surface"],
                                highlightthickness=0, bd=0)
                cir.pack()
                if es_sel:
                    cir.create_oval(1, 1, 39, 39, fill=C["accent"], outline="")
                    cir.create_text(20, 20, text=str(num_dia), fill=C["white"],
                                    font=F["small_bold"])
                else:
                    cir.create_text(20, 20, text=str(num_dia), fill=C["text_primary"],
                                    font=F["small_bold"])
                    if es_hoy:
                        cir.create_oval(1.5, 1.5, 38.5, 38.5, outline=C["accent"],
                                        width=2)
                color_marca = self.marcas.get(fecha.isoformat())
                if color_marca:
                    dot = tk.Canvas(cont, width=7, height=7, bg=C["surface"],
                                    highlightthickness=0, bd=0)
                    dot.pack(pady=(3, 0))
                    dot.create_oval(0, 0, 7, 7, fill=color_marca, outline="")
                for w in (celda, cont, cir):
                    w.bind("<Button-1>",
                           lambda e, d=num_dia: self._celda_click(d, anio, mes))
                self._celdas[(anio, mes, num_dia)] = {"c": celda}


class DashboardAdminApp:
    """Ventana principal del administrador.

    Agrupa los módulos Inicio, Asistencia, Reportes y Configuración en un
    área desplazable, con sidebar de navegación y topbar de acciones.
    """

    MODULOS = [
        ("inicio", "Inicio", "home"),
        ("asistencia", "Asistencia", "reloj"),
        ("reportes", "Reportes", "grafica"),
        ("config", "Configuraci\u00f3n", "config"),
    ]

    def __init__(self, usuario="admin", datos=None):
        """Configura el estado, crea la ventana y arranca carga y auto-refresco.

        Args:
            usuario (str): Nombre de usuario que inició sesión.
            datos (dict | None): Datos del usuario (nombre, correo, rol).
        """
        datos = datos or {}
        self.usuario = usuario
        self.nombre = datos.get("nombre") or "Administrador"
        self.correo = datos.get("correo", "")
        self.rol = datos.get("rol", "administrador")

        self.db = None
        self.firebase_listo = False
        self.hora_entrada = HORA_ENTRADA_DEFECTO
        self.hora_salida = HORA_SALIDA_DEFECTO
        self.nombre_empresa = "Empresa Qu\u00edmica"
        self.notificaciones = False

        self.usuarios = []
        self.marcaciones_hoy = []
        self.volver_login = False

        self.ventana = tk.Tk()
        self.ventana.title("Sistema de Asistencia - Administrador")
        configurar_icono(self.ventana)
        self.ventana.config(bg=C["bg_app"])
        self.ventana.overrideredirect(True)
        self._centrar_ventana(ANCHO, ALTO)
        self.ventana.minsize(ANCHO, ALTO)
        self.ventana.maxsize(ANCHO, ALTO)

        self.paginas = {}
        self.nav_botones = {}
        self._offset_x = 0
        self._offset_y = 0
        self._sidebar_colapsado = None
        self._cerrando = False
        self.pagina_actual = "inicio"
        self._alertas_pendientes = []
        self._alertas_popup = None
        self._after_reportes_id = None
        self._generando_reportes = False

        configure_styles(self.ventana)
        self._construir_ui()
        self._aplicar_responsividad()
        self.ventana.bind("<Configure>", lambda e: self._aplicar_responsividad())

        threading.Thread(target=self._inicializar_db, daemon=True).start()
        self._iniciar_polling()

        icons.forzar_taskbar(self.ventana)
        self.ventana.mainloop()

    def _centrar_ventana(self, w, h):
        """Centra la ventana en la pantalla con el tamaño indicado."""
        sw = self.ventana.winfo_screenwidth()
        sh = self.ventana.winfo_screenheight()
        x = (sw - w) // 2
        y = (sh - h) // 2
        self.ventana.geometry(f"{w}x{h}+{x}+{y}")

    # ------------------------------------------------------------------
    # BD
    # ------------------------------------------------------------------
    def _inicializar_db(self):
        """Conecta a Firestore y carga la configuración (se ejecuta en un hilo)."""
        try:
            self.db = get_db()
            self.firebase_listo = True
            self._cargar_config()
        except Exception as e:
            print(f"Error Firebase dashboard: {e}")
            self.ventana.after(0, self._mostrar_aviso_sin_db)

    def _mostrar_aviso_sin_db(self):
        """Muestra un aviso de error en la etiqueta de estado si Firebase falla."""
        if getattr(self, "_cerrando", False):
            return
        self.lbl_estado_inicio.config(
            text="No se pudo conectar a Firebase. Verifica config/firebase-key.json",
            fg=C["danger"])

    def _cargar_config(self):
        """Lee la configuración de empresa desde Firestore y la aplica."""
        try:
            doc = self.db.collection("config").document("empresa").get()
            if doc.exists:
                data = doc.to_dict()
                self.hora_entrada = data.get("hora_entrada", HORA_ENTRADA_DEFECTO)
                self.hora_salida = data.get("hora_salida", HORA_SALIDA_DEFECTO)
                self.nombre_empresa = data.get("nombre", "Empresa Qu\u00edmica")
                self.notificaciones = bool(data.get("notificaciones", False))
            self.ventana.after(0, self._poblar_config_ui)
        except Exception as e:
            print(e)
        self._cargar_usuarios()

    def _cargar_usuarios(self):
        """Carga todos los usuarios de Firestore y refresca las marcaciones de hoy."""
        try:
            docs = self.db.collection("usuarios").get()
            usuarios = []
            for d in docs:
                data = d.to_dict()
                data["__id"] = d.id
                usuarios.append(data)
            self.usuarios = usuarios
            self.ventana.after(0, lambda: self.lbl_estado_inicio.config(
                text=f"Conectado a Firestore. {len(self.usuarios)} usuario(s) cargado(s).",
                fg=C["success"]))
            self._cargar_marcaciones_hoy()
        except Exception as e:
            print(e)

    def _obtener_fecha_hoy(self):
        """Devuelve la fecha actual en formato YYYY-MM-DD."""
        return datetime.now().strftime("%Y-%m-%d")

    def _cargar_marcaciones_hoy(self):
        """Carga las marcaciones de hoy desde Firestore y refresca la interfaz."""
        try:
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", "==", self._obtener_fecha_hoy())).get()
            self.marcaciones_hoy = [d.to_dict() for d in docs]
            self.ventana.after(0, self._refrescar_inicio)
            self.ventana.after(100, self._refrescar_calendario)
        except Exception as e:
            print(e)

    def _trabajadores(self):
        """Lista de usuarios con rol distinto a administrador.

        Returns:
            list: Usuarios que son trabajadores.
        """
        return [u for u in self.usuarios
                if u.get("rol") != "administrador"]

    # ------------------------------------------------------------------
    # REFRESCO AUTOMATICO
    # ------------------------------------------------------------------
    def _iniciar_polling(self):
        """Programa el siguiente ciclo de auto-refresco después de POLL_INTERVALO ms."""
        if self._cerrando:
            return
        self._poll_after_id = self.ventana.after(POLL_INTERVALO, self._poll)

    def _detener_polling(self):
        """Cancela el ciclo de polling programado si existe.

        @return None: Detiene el auto-refresco antes de cerrar la ventana.
        """
        pid = getattr(self, "_poll_after_id", None)
        if pid is not None:
            try:
                self.ventana.after_cancel(pid)
            except Exception:
                pass
            self._poll_after_id = None

    def _poll(self):
        """Ciclo de polling: reprograma el siguiente y lanza la lectura en un hilo."""
        if self._cerrando:
            return
        self._iniciar_polling()
        if not getattr(self, "firebase_listo", False):
            return
        threading.Thread(target=self._poll_fetch, daemon=True).start()

    def _poll_fetch(self):
        """Lee las marcaciones de hoy desde Firestore (hilo) y aplica el refresco."""
        try:
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", "==",
                                          self._obtener_fecha_hoy())).get()
            marc = [d.to_dict() for d in docs]
            self.ventana.after(0, lambda m=marc: self._aplicar_refresco(m))
        except Exception as e:
            print(e)
            self._reintentar_carga_base()

    def _reintentar_carga_base(self):
        """Reintenta la carga de usuarios si la lectura inicial falló.

        Si el arranque del dashboard ocurrió sin lectura disponible (por
        ejemplo por cuota de lectura agotada), deja usuarios vacío. Este
        método relanza la carga en un hilo para que el siguiente ciclo de
        polling repueble los datos sin necesidad de reiniciar la aplicación.

        @return None: Relanza `_cargar_usuarios` en un hilo secundario.
        """
        if self._cerrando or not getattr(self, "firebase_listo", False):
            return
        if getattr(self, "usuarios", None):
            return
        threading.Thread(target=self._cargar_usuarios, daemon=True).start()

    def _aplicar_refresco(self, marc):
        """Aplica nuevos datos de marcaciones y refresca las páginas visibles."""
        if getattr(self, "_cerrando", False):
            return
        self.marcaciones_hoy = marc
        self._refrescar_inicio()
        self._refrescar_calendario()
        if self.pagina_actual == "asistencia" and self.db:
            threading.Thread(target=self._generar_asistencia, daemon=True).start()
        if self.pagina_actual == "reportes" and self.db:
            self._programar_generar_reportes()

    def _refrescar_manual(self):
        """Fuerza un refresco manual de datos desde el topbar."""
        self.lbl_estado_inicio.config(
            text="Actualizando datos...", fg=C["text_secondary"])
        self._poll()

    # ------------------------------------------------------------------
    # ESTRUCTURA
    # ------------------------------------------------------------------
    def _construir_ui(self):
        """Construye la topbar y el cuerpo de la ventana."""
        self._crear_topbar()
        self._crear_cuerpo()

    def _crear_topbar(self):
        """Crea la barra superior: título, acciones, campana de alertas y avatar."""
        topbar = tk.Frame(self.ventana, bg=C["bg_topbar"], height=ALTO_TOPBAR)
        topbar.pack(fill="x", side="top")
        topbar.pack_propagate(False)
        tk.Frame(topbar, bg=C["border"], height=1).pack(side="bottom", fill="x")

        icono_menu = icons.crear_icono("home", size=24, color=C["text_secondary"],
                                       bg=C["bg_topbar"], master=topbar)
        icono_menu.pack(side="left", padx=(14, 8), pady=7)
        tk.Label(topbar, text="SISTEMA DE ASISTENCIA",
                 bg=C["bg_topbar"], fg=C["text_secondary"],
                 font=F["small_bold"]).pack(side="left")

        btn_cerrar = tk.Label(topbar, text=" \u2715 ", bg=C["bg_topbar"],
                              fg=C["text_muted"], font=("Segoe UI", 10),
                              cursor="hand2")
        btn_cerrar.pack(side="right", padx=(0, 4))
        btn_cerrar.bind("<Enter>", lambda e: btn_cerrar.config(fg=C["accent"]))
        btn_cerrar.bind("<Leave>", lambda e: btn_cerrar.config(fg=C["text_muted"]))
        btn_cerrar.bind("<Button-1>", lambda e: self._cerrar_ventana_pura())

        btn_min = tk.Label(topbar, text=" \u2013 ", bg=C["bg_topbar"],
                           fg=C["text_muted"], font=("Segoe UI", 10),
                           cursor="hand2")
        btn_min.pack(side="right")
        btn_min.bind("<Enter>", lambda e: btn_min.config(fg=C["text_primary"]))
        btn_min.bind("<Leave>", lambda e: btn_min.config(fg=C["text_muted"]))
        btn_min.bind("<Button-1>", lambda e: self._minimizar())

        btn_refrescar = tk.Label(topbar, text=" \u21bb ", bg=C["bg_topbar"],
                                 fg=C["text_secondary"], font=("Segoe UI", 13),
                                 cursor="hand2")
        btn_refrescar.pack(side="right", padx=(0, 4), pady=6)
        btn_refrescar.bind("<Enter>",
                           lambda e: btn_refrescar.config(fg=C["accent"]))
        btn_refrescar.bind("<Leave>",
                           lambda e: btn_refrescar.config(fg=C["text_secondary"]))
        btn_refrescar.bind("<Button-1>", lambda e: self._refrescar_manual())

        cont_campana = tk.Frame(topbar, bg=C["bg_topbar"])
        cont_campana.pack(side="right", padx=(0, 12), pady=7)
        campana = icons.crear_icono("campana", size=20, color=C["text_secondary"],
                                    bg=C["bg_topbar"], master=cont_campana)
        campana.pack()
        campana.bind("<Button-1>", lambda e: self._toggle_popup_alertas())
        self._campana = campana
        self.badge_alertas = tk.Label(cont_campana, text="", bg=C["accent"],
                                      fg=C["white"], font=F["small_bold"],
                                      padx=4, pady=0, bd=0, highlightthickness=0)
        self._actualizar_badge_alertas()

        cont_avatar = tk.Frame(topbar, bg=C["bg_topbar"])
        cont_avatar.pack(side="right", padx=(8, 8), pady=5)
        avatar = tk.Canvas(cont_avatar, width=28, height=28, bg=C["bg_topbar"],
                           highlightthickness=0, bd=0)
        avatar.create_oval(1, 1, 27, 27, fill=C["surface_hover"], outline=C["border"])
        avatar.create_text(14, 14, text=self._iniciales(), fill=C["text_primary"],
                           font=F["small_bold"])
        avatar.pack(side="left")

        tk.Label(topbar, text="Administrador", bg=C["bg_topbar"],
                 fg=C["text_primary"], font=F["small_bold"]).pack(side="right")

        topbar.bind("<ButtonPress-1>", self._iniciar_arrastre)
        topbar.bind("<B1-Motion>", self._arrastrar)
        for hijo in topbar.winfo_children():
            self._bind_arrastre(hijo)
        campana.bind("<Button-1>", lambda e: self._toggle_popup_alertas())
        self.badge_alertas.bind("<Button-1>",
                                lambda e: self._toggle_popup_alertas())

    def _bind_arrastre(self, w):
        """Vincula los eventos de arrastre a un widget y a todos sus hijos."""
        for hijo in w.winfo_children():
            self._bind_arrastre(hijo)
        w.bind("<ButtonPress-1>", self._iniciar_arrastre)
        w.bind("<B1-Motion>", self._arrastrar)

    def _iniciales(self):
        """Devuelve las iniciales del nombre para el avatar del usuario."""
        partes = [p for p in self.nombre.split() if p]
        if len(partes) >= 2:
            return (partes[0][0] + partes[1][0]).upper()
        return (self.nombre[:2] or "AD").upper()

    def _cerrar_ventana_pura(self):
        """Marca el cierre y destruye la ventana del dashboard."""
        self._cerrando = True
        self._detener_polling()
        self.ventana.destroy()

    def _minimizar(self):
        """Minimiza la ventana desactivando temporalmente overrideredirect."""
        self.ventana.overrideredirect(False)
        self.ventana.iconify()
        self.ventana.after(600, lambda: (self.ventana.overrideredirect(True),
                                         icons.forzar_taskbar(self.ventana)))

    def _crear_cuerpo(self):
        """Crea el contenedor principal con la sidebar y el área de contenido."""
        self.cuerpo = tk.Frame(self.ventana, bg=C["bg_app"])
        self.cuerpo.pack(fill="both", expand=True)
        self.cuerpo.rowconfigure(0, weight=1)
        self.cuerpo.columnconfigure(1, weight=1)

        self._crear_sidebar(self.cuerpo)
        self._crear_contenido(self.cuerpo)

    def _crear_sidebar(self, padre):
        """Crea la barra lateral con branding, navegación y botón de salir."""
        self.sidebar = tk.Frame(padre, bg=C["bg_sidebar"], width=ANCHO_SIDEBAR)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_propagate(False)

        branding = tk.Frame(self.sidebar, bg=C["bg_sidebar"])
        branding.pack(fill="x", pady=(24, 4))
        tk.Label(branding, text="ASISTENCIA", bg=C["bg_sidebar"],
                 fg=C["text_primary"], font=("Segoe UI", 20, "bold"),
                 justify="center").pack()
        tk.Label(branding, text="ADMIN", bg=C["bg_sidebar"],
                 fg=C["accent"], font=("Segoe UI", 20, "bold"),
                 justify="center").pack()
        tk.Frame(self.sidebar, bg=C["accent"], height=2,
                 width=50).pack(pady=(8, 4))

        self._brand_children = branding.winfo_children()

        for clave, texto, icono in self.MODULOS:
            btn = NavButton(self.sidebar, texto, icono,
                            lambda c=clave: self._mostrar_modulo(c))
            btn.pack(fill="x", padx=14, pady=3)
            self.nav_botones[clave] = btn

        btn_usuarios = NavButton(self.sidebar, "Usuario", "usuarios",
                                 self._abrir_gestion_usuarios)
        btn_usuarios.pack(fill="x", padx=14, pady=3)

        btn_salir = tk.Button(
            self.sidebar, text="CERRAR SESI\u00d3N", bg=C["surface_raised"],
            fg=C["text_secondary"], activebackground=C["accent_hover"],
            activeforeground=C["white"], relief="flat", font=F["small_bold"],
            cursor="hand2", height=2, command=self._cerrar_sesion)
        btn_salir.pack(side="bottom", fill="x", padx=14, pady=16)

    def _crear_contenido(self, padre):
        """Crea el área desplazable y las cuatro páginas (módulos)."""
        cont_outer = tk.Frame(padre, bg=C["bg_app"])
        cont_outer.grid(row=0, column=1, sticky="nsew")
        cont_outer.rowconfigure(0, weight=1)
        cont_outer.columnconfigure(0, weight=1)

        self._scroll_canvas = tk.Canvas(cont_outer, bg=C["bg_app"],
                                        highlightthickness=0, bd=0)
        vsb = ttk.Scrollbar(cont_outer, orient="vertical", style="Dark.Scrollbar",
                            command=self._scroll_canvas.yview)
        self._scroll_canvas.configure(yscrollcommand=vsb.set)
        vsb.grid(row=0, column=1, sticky="ns")
        self._scroll_canvas.grid(row=0, column=0, sticky="nsew")

        self._scroll_frame = tk.Frame(self._scroll_canvas, bg=C["bg_app"])
        self._scroll_window = self._scroll_canvas.create_window(
            (0, 0), window=self._scroll_frame, anchor="nw")

        self._scroll_frame.bind(
            "<Configure>",
            lambda e: self._scroll_canvas.configure(
                scrollregion=self._scroll_canvas.bbox("all")))
        self._scroll_canvas.bind(
            "<Configure>", self._ajustar_scroll)
        self._scroll_canvas.bind("<MouseWheel>",
                                 lambda e: self._scroll_canvas.yview_scroll(
                                     int(-e.delta / 120), "units"))
        self._scroll_frame.bind("<MouseWheel>",
                                lambda e: self._scroll_canvas.yview_scroll(
                                    int(-e.delta / 120), "units"))

        self._crear_inicio()
        self._crear_asistencia()
        self._crear_reportes()
        self._crear_config()

        self._mostrar_modulo("inicio")

    def _ajustar_scroll(self, e):
        """Ajusta el área interna al ancho y alto del canvas de scroll."""
        self._scroll_canvas.itemconfigure(self._scroll_window, width=e.width)
        bbox = self._scroll_canvas.bbox("all")
        contenido_h = bbox[3] if bbox else 0
        self._scroll_canvas.itemconfigure(
            self._scroll_window, height=max(contenido_h, e.height))

    def _mostrar_modulo(self, clave):
        """Muestra el módulo indicado y actualiza la navegación lateral."""
        self.pagina_actual = clave
        for nombre, pagina in self.paginas.items():
            if nombre == clave:
                pagina.pack(in_=self._scroll_frame, fill="both", expand=True)
            else:
                pagina.pack_forget()
        self._scroll_canvas.yview_moveto(0)
        for nombre, btn in self.nav_botones.items():
            btn.set_activo(nombre == clave)
        self._cerrar_popup_alertas()
        if not self.firebase_listo:
            return
        if clave == "asistencia" and self.db:
            threading.Thread(target=self._generar_asistencia, daemon=True).start()
        elif clave == "reportes":
            self._lanzar_generar_reportes()

    def _aplicar_responsividad(self):
        """Ajusta el ancho de la sidebar y los paddings según el ancho de la ventana."""
        if getattr(self, "_cerrando", False):
            return
        try:
            if not self.ventana.winfo_exists():
                return
            ancho = self.ventana.winfo_width()
        except Exception:
            return
        try:
            if not self.sidebar.winfo_exists() or not self.cuerpo.winfo_exists():
                return
        except Exception:
            return
        contenido_inner = (ancho or ANCHO)
        if contenido_inner >= 1280:
            ancho_sidebar = 244
            self._pad_inner = 24
        elif contenido_inner >= 1000:
            ancho_sidebar = 200
            self._pad_inner = 18
        else:
            ancho_sidebar = 76
            self._pad_inner = 14

        if ancho_sidebar != self.sidebar.winfo_width():
            self.cuerpo.columnconfigure(0, minsize=ancho_sidebar)
            self.sidebar.config(width=ancho_sidebar)
        colapsado = contenido_inner < 1000
        if contenido_inner <= 10:
            colapsado = False
        if colapsado != self._sidebar_colapsado:
            self._sidebar_colapsado = colapsado
            for clave, btn in self.nav_botones.items():
                if colapsado:
                    btn._lbl.config(text="")
                else:
                    texto = dict((c, t) for c, t, _ in self.MODULOS).get(clave, "")
                    btn._lbl.config(text=texto)
                btn._redibujar()

    # ------------------------------------------------------------------
    # INICIO
    # ------------------------------------------------------------------
    def _crear_inicio(self):
        """Construye la página Inicio: KPI, estado, incidencias y actividad."""
        frame = tk.Frame(self._scroll_frame, bg=C["bg_app"])
        self.paginas["inicio"] = frame

        hoy = datetime.now().strftime("%d/%m/%Y")
        PageHeader(frame, f"Buenos d\u00edas, {self.nombre}",
                   f"Revisa el estado de asistencia de tu equipo | {hoy}"
                   ).pack(fill="x", padx=24, pady=(14, 2))

        self.lbl_estado_inicio = tk.Label(
            frame, text="Conectando con Firebase...", bg=C["bg_app"],
            fg=C["text_secondary"], font=F["small_gris"] if hasattr(F, "small_gris")
            else F["small"], anchor="w", justify="left")
        self.lbl_estado_inicio.pack(anchor="w", padx=24)

        fila_kpis = tk.Frame(frame, bg=C["bg_app"])
        fila_kpis.pack(fill="x", padx=24, pady=(8, 4))

        self.kpis = {}
        self._metric_cards = []
        defs = [
            ("kpi_presentes", "Presentes", "success", "reloj"),
            ("kpi_atrasados", "Atrasados", "warning", "exclamacion"),
            ("kpi_ausentes", "Ausentes", "danger", "cerrar"),
        ]
        for i, (clave, texto, color, icono) in enumerate(defs):
            card = MetricCard(fila_kpis, texto, "--", "de hoy",
                              color=C[color], icono=icono, tamano=(None, 116))
            card.pack(side="left", fill="both", expand=True,
                      padx=(0 if i == 0 else 8, 0 if i == 2 else 8))
            self.kpis[clave] = card._dibujar
            self._metric_cards.append(card)

        self.medio = tk.Frame(frame, bg=C["bg_app"])
        self.medio.pack(fill="both", expand=True, padx=24, pady=(2, 2))
        self.medio.columnconfigure(0, weight=3)
        self.medio.columnconfigure(1, weight=2)
        self.medio.rowconfigure(0, weight=1)

        self._crear_card_estado(self.medio)
        self._crear_card_incidencias(self.medio)
        self._crear_card_actividad(frame)

    def _actualizar_kpis(self, valores):
        """Actualiza las tarjetas de KPI con los valores dados (clave, valor, sub)."""
        for clave, valor, sub in valores:
            for card in self._metric_cards:
                if card.titulo.lower().replace(" ", "_") == clave.replace("kpi_", ""):
                    card.valor = str(valor)
                    card.subtitulo = sub
                    card._dibujar()

    def _crear_card_estado(self, padre):
        """Crea el card ESTADO DE ASISTENCIA con las barras horizontales."""
        card = SurfaceCard(padre)
        card.grid(row=0, column=0, sticky="nsew", padx=(0, 4), pady=4)

        tk.Label(card, text="ESTADO DE ASISTENCIA", bg=C["surface"],
                 fg=C["text_secondary"], font=F["small_bold"]).pack(
            anchor="w", padx=16, pady=(14, 2))

        cont = tk.Frame(card, bg=C["surface"])
        cont.pack(fill="both", expand=True, padx=8, pady=(0, 10))

        self.cont_barras = tk.Frame(cont, bg=C["surface"])
        self.cont_barras.pack(fill="both", expand=True)

    def _crear_card_incidencias(self, padre):
        """Crea el card de incidencias del día."""
        card = SurfaceCard(padre)
        card.grid(row=0, column=1, sticky="nsew", padx=(4, 0), pady=4)

        tk.Label(card, text="INCIDENCIAS DEL D\u00cdA", bg=C["surface"],
                 fg=C["warning"], font=F["small_bold"]).pack(
            anchor="w", padx=16, pady=(14, 2))

        self.cont_incidencias = tk.Frame(card, bg=C["surface"])
        self.cont_incidencias.pack(fill="both", expand=True, padx=10, pady=(0, 12))

    def _crear_card_actividad(self, frame):
        """Crea el card de actividad reciente (últimas marcaciones)."""
        card = SurfaceCard(frame, padding=10)
        card.pack(fill="x", padx=24, pady=(2, 14))

        tk.Label(card, text="ACTIVIDAD RECIENTE", bg=C["surface"],
                 fg=C["text_secondary"], font=F["small_bold"]).pack(
            anchor="w", padx=4, pady=(4, 2))
        cont = tk.Frame(card, bg=C["surface"])
        cont.pack(fill="x", padx=4)
        linea = tk.Frame(cont, bg=C["border"], width=2)
        linea.pack(side="left", fill="y", padx=(13, 10), pady=2)
        self.lista_actividad = tk.Frame(cont, bg=C["surface"])
        self.lista_actividad.pack(side="left", fill="x", expand=True)

    # ------------------------------------------------------------------
    # INICIO - REFRESCO
    # ------------------------------------------------------------------
    def _refrescar_inicio(self):
        """Recalcula KPIs, barras, incidencias, actividad y alertas con datos del día."""
        marc = self.marcaciones_hoy
        trabajadores = self._trabajadores()
        nombres_trabajo = {u.get("usuario") for u in trabajadores}

        total = len(trabajadores)
        presentes = {}
        for m in marc:
            if m.get("usuario") not in nombres_trabajo:
                continue
            presentes.setdefault(m.get("usuario"), {"entrada": None, "salida": None})
            if m.get("tipo") == "entrada":
                presentes[m["usuario"]]["entrada"] = m.get("hora")
            elif m.get("tipo") == "salida":
                presentes[m["usuario"]]["salida"] = m.get("hora")

        n_presentes = len(presentes)
        n_atrasados = sum(1 for d in presentes.values()
                          if d["entrada"] and d["entrada"] > self.hora_entrada)
        n_salida_ant = sum(1 for d in presentes.values()
                           if d["salida"] and d["salida"] < self.hora_salida)

        ausentes = nombres_trabajo - set(presentes)
        n_ausentes = len(ausentes)

        self._actualizar_kpis([
            ("kpi_presentes", n_presentes, "de hoy"),
            ("kpi_atrasados", n_atrasados, "de hoy"),
            ("kpi_ausentes", n_ausentes, "de hoy"),
        ])

        self._dibujar_barras(total, n_presentes, n_atrasados, n_ausentes)
        self._poblar_incidencias(presentes, ausentes, marc)
        self._poblar_actividad(marc, trabajadores)
        threading.Thread(target=self._registrar_alertas_inasistencia,
                         args=(ausentes,), daemon=True).start()
        threading.Thread(target=self._cargar_alertas_pendientes,
                         daemon=True).start()
        self._marcar_ultima_actualizacion(n_presentes)

    def _marcar_ultima_actualizacion(self, n_presentes=0):
        """Actualiza la etiqueta de estado con hora, conteos y trabajadores."""
        hora = datetime.now().strftime("%H:%M:%S")
        self.lbl_estado_inicio.config(
            text=f"Conectado a Firestore \u00b7 {len(self._trabajadores())} "
                 f"trabajadores \u00b7 Actualizado {hora} \u00b7 {n_presentes} "
                 f"presente(s) registrado(s).",
            fg=C["success"])

    def _dibujar_barras(self, total, presentes, atrasados, ausentes):
        """Dibuja las barras horizontales de Presentes, Atrasados y Ausentes.

        Args:
            total (int): Total de trabajadores.
            presentes (int): Trabajadores con registro de hoy.
            atrasados (int): Trabajadores con entrada después de la hora límite.
            ausentes (int): Trabajadores sin registro.
        """
        for w in self.cont_barras.winfo_children():
            w.destroy()

        segmentos = [
            ("Presentes", presentes, C["success"]),
            ("Atrasados", atrasados, C["warning"]),
            ("Ausentes", ausentes, C["danger"]),
        ]

        if total == 0:
            tk.Label(self.cont_barras, text="Sin datos", bg=C["surface"],
                     fg=C["text_muted"], font=F["body_bold"]).pack(
                anchor="center", pady=18)
            return

        alto_track = 14
        for nombre, valor, color in segmentos:
            fila = tk.Frame(self.cont_barras, bg=C["surface"])
            fila.pack(fill="x", pady=6)
            fila.columnconfigure(1, weight=1)

            izq = tk.Frame(fila, bg=C["surface"])
            izq.grid(row=0, column=0, sticky="w")
            dot = tk.Canvas(izq, width=12, height=12, bg=C["surface"],
                            highlightthickness=0, bd=0)
            dot.create_oval(1, 1, 11, 11, fill=color, outline="")
            dot.pack(side="left", padx=(0, 8))
            tk.Label(izq, text=nombre, bg=C["surface"], fg=C["text_secondary"],
                     font=F["small"]).pack(side="left")

            track = tk.Frame(fila, bg=C["input_bg"], height=alto_track)
            track.grid(row=0, column=1, sticky="ew", padx=(12, 0))
            track.pack_propagate(False)

            frac = valor / total if total else 0.0
            if frac > 0:
                barra = tk.Frame(track, bg=color)
                barra.place(x=0, y=0, relwidth=frac, relheight=1)

            tk.Label(fila, text=str(valor), bg=C["surface"],
                     fg=C["text_primary"], font=F["body_bold"]).grid(
                row=0, column=2, sticky="e", padx=(12, 0))

    def _poblar_incidencias(self, presentes, ausentes, marc):
        """Llena el card de incidencias con inasistencias, atrasos y salidas anticipadas."""
        for w in self.cont_incidencias.winfo_children():
            w.destroy()

        trabajo_ids = {u.get("usuario") for u in self._trabajadores()}
        usuarios_nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                            for u in self.usuarios}

        filas = []
        for u in sorted(ausentes):
            filas.append((usuarios_nombres.get(u, u), "Inasistencia", "Sin registros",
                          C["danger"]))
        combs = {}
        for m in marc:
            if m.get("usuario") not in trabajo_ids:
                continue
            combs.setdefault(m.get("usuario"), {})
            if m.get("tipo") == "entrada":
                combs[m["usuario"]]["entrada"] = m.get("hora")
            elif m.get("tipo") == "salida":
                combs[m["usuario"]]["salida"] = m.get("hora")
        for u, d in combs.items():
            nombre = usuarios_nombres.get(u, u)
            if d.get("entrada") and d["entrada"] > self.hora_entrada:
                filas.append((nombre, "Atraso", d["entrada"], C["warning"]))
            if d.get("salida") and d["salida"] < self.hora_salida:
                filas.append((nombre, "Salida anticipada", d["salida"], C["warning"]))

        if not filas:
            EmptyState(self.cont_incidencias, "Sin incidencias",
                       "Nadie present\u00f3 atrasos ni ausencias hoy",
                       icono="check", color=C["success"]).pack(
                fill="both", expand=True)
            return

        for nombre, tipo, detalle, color in filas[:8]:
            fila = tk.Frame(self.cont_incidencias, bg=C["surface"])
            fila.pack(fill="x", pady=3)
            fila.config(height=64)
            fila.pack_propagate(False)

            avatar = tk.Canvas(fila, width=42, height=42, bg=C["surface"],
                               highlightthickness=0, bd=0)
            avatar.create_oval(1, 1, 41, 41, fill=C["surface_raised"],
                               outline=C["border"])
            inits = "".join(p[0] for p in nombre.split()[:2]).upper() or "?"
            avatar.create_text(21, 21, text=inits, fill=C["text_secondary"],
                               font=F["small_bold"])
            avatar.pack(side="left", padx=(12, 10), pady=11)

            medio = tk.Frame(fila, bg=C["surface"])
            medio.pack(side="left", fill="both", expand=True, pady=9)
            medio.columnconfigure(0, weight=0)
            medio.columnconfigure(1, weight=1)
            tk.Label(medio, text=nombre, bg=C["surface"], fg=C["text_primary"],
                     font=F["body_bold"], anchor="w").grid(
                row=0, column=0, columnspan=2, sticky="w")
            StatusBadge(medio, tipo,
                        "danger" if color == C["danger"] else "warning").grid(
                row=1, column=0, sticky="w", pady=(4, 0))
            tk.Label(medio, text=detalle, bg=C["surface"], fg=C["text_muted"],
                     font=F["small"], anchor="w").grid(
                row=1, column=1, sticky="w", padx=(8, 0), pady=(6, 0))

            tk.Button(fila, text="Ver detalle", bg=C["surface_raised"],
                      fg=C["text_secondary"], activebackground=C["surface_hover"],
                      activeforeground=C["text_primary"], relief="flat",
                      font=F["small"], cursor="hand2", padx=8,
                      command=lambda n=nombre: self._ver_detalle_incidencia(n))\
                .pack(side="right", padx=(4, 12), pady=18)

    def _ver_detalle_incidencia(self, nombre_usuario):
        """Abre una ventana auxiliar con el detalle de una incidencia."""
        from tkinter import Toplevel
        ventana = Toplevel(self.ventana)
        ventana.title("Detalle de incidencia")
        ventana.config(bg=C["surface"])
        ventana.geometry("460x300")
        ventana.resizable(False, False)
        tk.Label(ventana, text=f"Incidencia de {nombre_usuario}",
                 bg=C["surface"], fg=C["text_primary"], font=F["card_title"]).pack(
            padx=20, pady=(20, 10), anchor="w")
        tk.Frame(ventana, bg=C["border"], height=1).pack(fill="x", padx=20)
        tk.Label(ventana,
                 text="Esta incidencia se gener\u00f3 hoy.\n"
                      "El estado del trabajador figura en la pesta\u00f1a de Asistencia.\n\n"
                      "Revisa el reporte de inasistencias para m\u00e1s contexto.",
                 bg=C["surface"], fg=C["text_secondary"], font=F["body"],
                 justify="left").pack(padx=20, pady=(14, 10), anchor="w")
        SecondaryButton(ventana, "CERRAR", command=ventana.destroy).pack(
            side="bottom", anchor="e", padx=20, pady=14)

    def _poblar_actividad(self, marc, usuarios):
        """Muestra los últimos eventos de entrada/salida registrados hoy."""
        for w in self.lista_actividad.winfo_children():
            w.destroy()
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in usuarios}

        eventos = []
        for m in marc:
            accion = "registr\u00f3 entrada" if m.get("tipo") == "entrada" else "registr\u00f3 salida"
            nm = nombres.get(m.get("usuario"), m.get("usuario") or "?")
            color = C["success"] if m.get("tipo") == "entrada" else C["info"]
            eventos.append((nm, accion, m.get("hora", ""), color))

        if not eventos:
            tk.Label(self.lista_actividad, text="No hay registros de hoy todav\u00eda.",
                     bg=C["surface"], fg=C["text_muted"], font=F["small"]).pack(
                anchor="w", pady=6)
            return

        for nm, accion, hora, color in eventos[-3:]:
            evento = tk.Frame(self.lista_actividad, bg=C["surface"])
            evento.pack(fill="x", pady=4)
            circulo = tk.Canvas(evento, width=30, height=30, bg=C["surface"],
                                highlightthickness=0, bd=0)
            circulo.create_oval(1, 1, 29, 29, fill=C["surface_raised"],
                                outline=color, width=2)
            circulo.create_text(15, 15, text="\u2713" if accion.startswith("registr\u00f3 entrada")
                                else "\u203a", fill=color, font=F["small_bold"])
            circulo.pack(side="left", padx=(0, 10))
            tk.Label(evento, text=f"{nm} {accion}", bg=C["surface"],
                     fg=C["text_primary"], font=F["body"], anchor="w").pack(
                side="left")
            tk.Label(evento, text=f"{hora}", bg=C["surface"], fg=C["text_muted"],
                     font=F["small"]).pack(side="right", padx=(8, 0))

    def _registrar_alertas_inasistencia(self, ausentes):
        """Reconcilia las alertas de inasistencia de hoy con los ausentes reales.

        Elimina las alertas de inasistencia de hoy para trabajadores que sí
        tienen marcación (falsos positivos generados cuando arranca sin datos)
        y crea las que faltan para los ausentes reales.
        """
        try:
            if not self.db:
                return
            if not getattr(self, "marcaciones_hoy", []):
                return
            hoy = self._obtener_fecha_hoy()
            ausentes = set(ausentes)
            existentes = set()
            a_eliminar = []
            for d in self.db.collection("alertas") \
                    .where(filter=FieldFilter("fecha", "==", hoy)).get():
                datos = d.to_dict()
                if datos.get("tipo") != Alerta.TIPO_INASISTENCIA:
                    continue
                existentes.add(datos.get("usuario"))
                if datos.get("usuario") not in ausentes:
                    a_eliminar.append(d.reference)
            for ref in a_eliminar:
                ref.delete()
            for u in ausentes - existentes:
                Alerta.crear(self.db, Alerta.TIPO_INASISTENCIA, u, hoy)
        except Exception as e:
            print(e)

    def _cargar_alertas_pendientes(self):
        """Carga, ordena y formatea las alertas pendientes para el badge."""
        try:
            if not self.db:
                return
            hoy = self._obtener_fecha_hoy()
            alertas = Alerta.listar(self.db, estado=Alerta.ESTADO_PENDIENTE,
                                    fecha_inicio=hoy)
            orden = {Alerta.TIPO_ATRASO: 0, Alerta.TIPO_SALIDA_ANTICIPADA: 1,
                     Alerta.TIPO_INASISTENCIA: 2}
            alertas.sort(key=lambda a: (orden.get(a.get("tipo"), 3),
                                        a.get("fecha", "")))
            etiquetas = {
                Alerta.TIPO_ATRASO: "Atraso",
                Alerta.TIPO_SALIDA_ANTICIPADA: "Salida anticipada",
                Alerta.TIPO_INASISTENCIA: "Inasistencia",
            }
            nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                       for u in self.usuarios}
            self._alertas_pendientes = [
                (nombres.get(a.get("usuario"), a.get("usuario") or "?"),
                 etiquetas.get(a.get("tipo"), a.get("tipo")),
                 a.get("hora") or a.get("fecha") or "", a.get("_id"))
                for a in alertas
            ]
            self.ventana.after(0, lambda: (
                self._actualizar_badge_alertas(),
                self._reconstruir_popup_alertas()))
        except Exception as e:
            print(e)

    def _actualizar_badge_alertas(self):
        """Muestra u oculta el badge con el conteo de alertas pendientes."""
        n = len(getattr(self, "_alertas_pendientes", []) or [])
        if n <= 0:
            self.badge_alertas.place_forget()
        else:
            self.badge_alertas.config(text=str(n if n < 99 else "99+"))
            self.badge_alertas.place(relx=0.85, rely=0.05, anchor="ne")

    def _toggle_popup_alertas(self):
        """Alterna la apertura de la campana de alertas pendientes."""
        pop = getattr(self, "_alertas_popup", None)
        if pop is not None and pop.winfo_exists():
            self._cerrar_popup_alertas()
            return
        self._abrir_popup_alertas()
        threading.Thread(target=self._cargar_alertas_pendientes,
                         daemon=True).start()

    def _cerrar_popup_alertas(self):
        """Destruye el desplegable de la campana si está abierto."""
        pop = getattr(self, "_alertas_popup", None)
        if pop is not None:
            try:
                pop.destroy()
            except Exception:
                pass
        self._alertas_popup = None

    def _abrir_popup_alertas(self):
        """Abre el desplegable de alertas pendientes debajo de la campana."""
        self._cerrar_popup_alertas()
        if not getattr(self, "firebase_listo", False):
            self._mostrar_toast("Sin conexión a la base de datos", ok=False)
            return
        pop = tk.Toplevel(self.ventana)
        pop.overrideredirect(True)
        pop.configure(bg=C["surface"], highlightthickness=1,
                      highlightbackground=C["border"])
        try:
            pop.attributes("-topmost", True)
        except tk.TclError:
            pass
        self._alertas_popup = pop
        self._reconstruir_popup_alertas()

    def _reconstruir_popup_alertas(self):
        """Pinta el contenido del desplegable de alertas (reconstruible)."""
        pop = getattr(self, "_alertas_popup", None)
        if pop is None or not pop.winfo_exists():
            return
        for hijo in pop.winfo_children():
            hijo.destroy()
        pendientes = list(getattr(self, "_alertas_pendientes", []) or [])
        ancho = 320

        cab = tk.Frame(pop, bg=C["surface_raised"], padx=10, pady=6)
        cab.pack(fill="x")
        tk.Label(cab, text="Alertas pendientes", bg=C["surface_raised"],
                 fg=C["text_primary"], font=F["small_bold"]).pack(side="left")
        tk.Label(cab, text=f"({len(pendientes)})", bg=C["surface_raised"],
                 fg=C["accent"], font=F["small_bold"]).pack(side="right")

        if not pendientes:
            tk.Label(pop, text="Sin alertas pendientes", bg=C["surface"],
                     fg=C["text_muted"], font=F["small"], pady=16).pack(fill="x")

        for nombre, etiqueta, hora, alerta_id in pendientes[:8]:
            fila = tk.Frame(pop, bg=C["surface"], padx=10, pady=5)
            fila.pack(fill="x")
            izq = tk.Frame(fila, bg=C["surface"])
            izq.pack(side="left", fill="x", expand=True)
            tk.Label(izq, text=f"{nombre} · {etiqueta}", bg=C["surface"],
                     fg=C["text_primary"], font=F["small_bold"]).pack(anchor="w")
            tk.Label(izq, text=hora, bg=C["surface"], fg=C["text_muted"],
                     font=F["small"]).pack(anchor="w")
            lnk = tk.Label(fila, text="Leída", bg=C["surface"], fg=C["accent"],
                           font=F["small_bold"], cursor="hand2")
            lnk.pack(side="right", padx=(6, 4))
            lnk.bind("<Button-1>",
                     lambda e, i=alerta_id: self._marcar_alerta_popup([i]))
            lnk.bind("<Enter>", lambda e, l=lnk: l.config(fg=C["text_primary"]))
            lnk.bind("<Leave>", lambda e, l=lnk: l.config(fg=C["accent"]))

        if len(pendientes) > 8:
            tk.Label(pop, text=f"+{len(pendientes) - 8} más...", bg=C["surface"],
                     fg=C["text_muted"], font=F["small"]).pack(fill="x")

        pie = tk.Frame(pop, bg=C["surface_raised"], padx=10, pady=6)
        pie.pack(fill="x")
        if pendientes:
            btn_todas = tk.Label(pie, text="Marcar todas", bg=C["surface_raised"],
                                 fg=C["accent"], font=F["small_bold"],
                                 cursor="hand2")
            btn_todas.pack(side="left")
            btn_todas.bind("<Button-1>", lambda e: self._marcar_alerta_popup(
                [a[3] for a in pendientes]))
        btn_ver = tk.Label(pie, text="Ver en Reportes", bg=C["surface_raised"],
                           fg=C["accent"], font=F["small_bold"], cursor="hand2")
        btn_ver.pack(side="right")
        btn_ver.bind("<Button-1>", lambda e: self._ir_a_alertas_reportes())

        pop.update_idletasks()
        x = self._campana.winfo_rootx() - ancho + 24
        y = self._campana.winfo_rooty() + self._campana.winfo_height() + 4
        pop.geometry(f"{ancho}x{pop.winfo_reqheight()}+{x}+{y}")

    def _marcar_alerta_popup(self, ids):
        """Marca como leídas alertas desde la campana, en un hilo."""
        threading.Thread(target=self._marcar_alertas_leidas,
                         args=(list(ids),), daemon=True).start()

    def _marcar_alertas_leidas(self, ids):
        """Marca alertas en Firestore y refresca campana, badge y reportes.

        Se ejecuta en un hilo para no bloquear la interfaz de Tk.
        """
        for alerta_id in ids:
            try:
                Alerta.marcar_leida(self.db, alerta_id)
            except Exception as e:
                print(e)
        self._cargar_alertas_pendientes()
        self._generar_alertas(self.entry_desde.get().strip(),
                              self.entry_hasta.get().strip())
        self.ventana.after(0, lambda: (
            self._actualizar_metricas(),
            self._actualizar_titulos_tabs(),
            self._aplicar_filtro_reporte(),
            self._mostrar_toast(
                f"{len(ids)} alerta(s) marcada(s) como leída(s)")))

    def _ir_a_alertas_reportes(self):
        """Cierra la campana y abre la pestaña Alertas de la página Reportes."""
        self._cerrar_popup_alertas()
        self._mostrar_modulo("reportes")
        if hasattr(self, "notebook_reportes"):
            self.notebook_reportes.select(3)

    # ------------------------------------------------------------------
    # ASISTENCIA
    # ------------------------------------------------------------------
    def _crear_asistencia(self):
        """Construye la página Asistencia: calendario, búsqueda, filtros y tabla."""
        frame = tk.Frame(self._scroll_frame, bg=C["bg_app"])
        self.paginas["asistencia"] = frame

        PageHeader(frame, "Asistencia",
                   "Registros diarios de entrada y salida."
                   ).pack(fill="x", padx=24, pady=(20, 10))

        cont = tk.Frame(frame, bg=C["bg_app"])
        cont.pack(fill="both", expand=True, padx=24, pady=(0, 20))
        cont.columnconfigure(1, weight=1)
        cont.rowconfigure(0, weight=1)

        self.fecha_seleccionada = self._obtener_fecha_hoy()
        marcas = self._marcas_calendario()
        self.calendario = CalendarWidget(
            cont, seleccionada=datetime.strptime(self.fecha_seleccionada, "%Y-%m-%d").date(),
            on_seleccion=self._seleccionar_fecha, marcas=marcas)
        self.calendario.grid(row=0, column=0, sticky="nsw", padx=(0, 16))

        derecha = tk.Frame(cont, bg=C["bg_app"])
        derecha.grid(row=0, column=1, sticky="nsew")
        derecha.columnconfigure(0, weight=1)
        derecha.rowconfigure(1, weight=1)

        barra = tk.Frame(derecha, bg=C["bg_app"])
        barra.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        self.entry_buscar_asistencia = tk.Entry(
            barra, bg=C["input_bg"], fg=C["text_primary"], relief="flat",
            insertbackground=C["text_primary"], font=F["body"],
            highlightthickness=1, highlightbackground=C["border"],
            highlightcolor=C["border_active"])
        self.entry_buscar_asistencia.insert(0, "Buscar trabajador...")
        self.entry_buscar_asistencia.pack(side="left", fill="x", expand=True,
                                          padx=(0, 8), ipady=9)
        self.entry_buscar_asistencia.bind("<FocusIn>", self._placeholder_quitar)
        self.entry_buscar_asistencia.bind("<FocusOut>", self._placeholder_volver)
        self.entry_buscar_asistencia.bind(
            "<KeyRelease>", lambda e: self._aplicar_filtros_asistencia())

        self.estados_posibles = [
            "Todos", "Normal", "Atrasado", "Salida anticipada",
            "Inasistencia", "Solo salida", "Pendiente de salida",
            "Atrasado y salida anticipada",
        ]
        self.combo_estado = ttk.Combobox(
            barra, values=self.estados_posibles, state="readonly",
            style="Dark.TCombobox", font=F["body"], width=24)
        self.combo_estado.set("Todos")
        self.combo_estado.pack(side="left", padx=(0, 8), ipady=3)
        self.combo_estado.bind("<<ComboboxSelected>>",
                               lambda e: self._aplicar_filtros_asistencia())

        self.lbl_asistencia = tk.Label(barra, text="", bg=C["bg_app"],
                                       fg=C["text_secondary"], font=F["small"])
        self.lbl_asistencia.pack(side="left", padx=8)

        cont_tabla = tk.Frame(derecha, bg=C["surface"], highlightthickness=1,
                              highlightbackground=C["border"])
        cont_tabla.grid(row=1, column=0, sticky="nsew")

        cols = ("trabajador", "fecha", "entrada", "salida", "estado")
        self.tabla_asistencia = ttk.Treeview(
            cont_tabla, columns=cols, show="headings", selectmode="browse",
            style="Dark.Treeview")
        for c, titulo, ancho in [
            ("trabajador", "Trabajador", 140),
            ("fecha", "Fecha", 84),
            ("entrada", "Entrada", 84),
            ("salida", "Salida", 84),
            ("estado", "Estado", 140),
        ]:
            self.tabla_asistencia.heading(c, text=titulo)
            self.tabla_asistencia.column(c, width=ancho)
        self._configurar_tags_estado(self.tabla_asistencia, 4)
        self.tabla_asistencia.pack(side="left", fill="both", expand=True)
        scroll = ttk.Scrollbar(cont_tabla, orient="vertical",
                               command=self.tabla_asistencia.yview,
                               style="Dark.Scrollbar")
        scroll.pack(side="right", fill="y")
        self.tabla_asistencia.configure(yscrollcommand=scroll.set)

    def _marcas_calendario(self):
        """Devuelve un color por fecha según la proporción de presentes.

        Returns:
            dict: Fecha (YYYY-MM-DD) como clave y color como valor.
        """
        marcas = {}
        try:
            if self.db:
                base = getattr(getattr(self, "calendario", None), "mostrando",
                           datetime.now().date())
                inicio = base.replace(day=1)
                if inicio.month == 12:
                    fin = inicio.replace(year=inicio.year + 1, month=1,
                                         day=1) - timedelta(days=1)
                else:
                    fin = inicio.replace(month=inicio.month + 1,
                                         day=1) - timedelta(days=1)
                docs = self.db.collection("marcaciones") \
                    .where(filter=FieldFilter("fecha", ">=",
                                              inicio.strftime("%Y-%m-%d"))) \
                    .where(filter=FieldFilter("fecha", "<=",
                                              fin.strftime("%Y-%m-%d"))).get()
                por_fecha = {}
                for d in docs:
                    m = d.to_dict()
                    por_fecha.setdefault(m.get("fecha"), set()).add(m.get("usuario"))
                nombres = {u.get("usuario") for u in self._trabajadores()}
                for fecha, presentes in por_fecha.items():
                    if len(presentes) == len(nombres) and nombres:
                        marcas[fecha] = C["success"]
                    elif not presentes:
                        marcas[fecha] = C["danger"]
                    else:
                        marcas[fecha] = C["warning"]
        except Exception as e:
            print(e)
        return marcas

    def _refrescar_calendario(self):
        """Recalcula las marcas del calendario y redibuja su cuadrícula."""
        calendario = getattr(self, "calendario", None)
        if calendario is None:
            return

        def _obtener_marcas():
            """Consulta las marcas en un hilo y las aplica al calendario."""
            try:
                marcas = self._marcas_calendario()
                self.ventana.after(
                    0, lambda: (setattr(calendario, "marcas", marcas),
                                calendario._llenar_cuadricula()))
            except Exception as e:
                print(e)

        threading.Thread(target=_obtener_marcas, daemon=True).start()

    def _seleccionar_fecha(self, fecha_str):
        """Guarda la fecha elegida y lanza la generación de la tabla en un hilo."""
        self.fecha_seleccionada = fecha_str
        threading.Thread(target=self._generar_asistencia, daemon=True).start()

    def _placeholder_quitar(self, e):
        """Borra el texto de ejemplo del buscador al recibir el foco."""
        if self.entry_buscar_asistencia.get() == "Buscar trabajador...":
            self.entry_buscar_asistencia.delete(0, "end")
            self.entry_buscar_asistencia.config(fg=C["text_primary"])

    def _placeholder_volver(self, e):
        """Reinserta el texto de ejemplo si el buscador queda vacío."""
        if not self.entry_buscar_asistencia.get().strip():
            self.entry_buscar_asistencia.insert(0, "Buscar trabajador...")
            self.entry_buscar_asistencia.config(fg=C["text_muted"])

    def _configurar_tags_estado(self, tabla, col_estado):
        """Configura los colores de estado de una tabla y guarda su índice."""
        self._tag_col = col_estado
        tabla.tag_configure("normal", foreground=C["success"])
        tabla.tag_configure("atraso", foreground=C["warning"])
        tabla.tag_configure("inasistencia", foreground=C["danger"])
        tabla.tag_configure("pendiente", foreground=C["text_secondary"])

    def _poblar_tabla_estado(self, tabla, filas):
        """Rellena una tabla con las filas de estado y colorea la columna de estado."""
        for item in tabla.get_children():
            tabla.delete(item)
        if not filas:
            tabla.insert("", "end", values=("Sin coincidencias", "", "", "", ""))
            return
        for f in filas:
            estado = f[self._tag_col]
            tag = "normal"
            if "Inasistencia" in estado:
                tag = "inasistencia"
            elif "Atrasado" in estado:
                tag = "atraso"
            elif "Solo salida" in estado or "Pendiente de salida" in estado:
                tag = "pendiente"
            tabla.insert("", "end", values=f, tags=(tag,))

    def _aplicar_filtros_asistencia(self):
        """Filtra las filas de asistencia por texto y estado, y repuebla la tabla."""
        if not hasattr(self, "_filas_asistencia"):
            return
        busqueda = self.entry_buscar_asistencia.get().strip().lower()
        if busqueda == "buscar trabajador...":
            busqueda = ""
        estado = self.combo_estado.get()
        filas = []
        for f in self._filas_asistencia:
            if busqueda and busqueda not in f[0].lower():
                continue
            if estado != "Todos" and f[4] != estado:
                continue
            filas.append(f)
        self._poblar_tabla_estado(self.tabla_asistencia, filas)

    def _generar_asistencia(self):
        """Consulta Firestore y arma las filas de estado de todos los trabajadores."""
        fecha = self.fecha_seleccionada
        if fecha in FERIADOS:
            self._filas_asistencia = []
            self.ventana.after(0, lambda: (
                self._aplicar_filtros_asistencia(),
                self.lbl_asistencia.config(
                    text="D\u00eda no laboral (feriado).", fg=C["text_secondary"])))
            return
        try:
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", "==", fecha)).get()
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_asistencia.config(
                text=str(e), fg=C["danger"]))
            return

        by_user = {}
        for d in docs:
            m = d.to_dict()
            by_user.setdefault(m.get("usuario"), {"entrada": None, "salida": None})
            if m.get("tipo") == "entrada":
                by_user[m["usuario"]]["entrada"] = m.get("hora")
            elif m.get("tipo") == "salida":
                by_user[m["usuario"]]["salida"] = m.get("hora")

        trabajadores = self._trabajadores()
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in trabajadores}

        filas = []
        for u in trabajadores:
            d = by_user.get(u.get("usuario"), {"entrada": None, "salida": None})
            filas.append((
                nombres.get(u.get("usuario"), u.get("usuario")),
                fecha,
                d["entrada"] or "Sin registro",
                d["salida"] or "Sin registro",
                self._calcular_estado(d["entrada"], d["salida"]),
            ))
        filas.sort(key=lambda r: (r[4] == "Inasistencia", r[0].lower()))
        self._filas_asistencia = filas
        self.ventana.after(0, lambda: (
            self._aplicar_filtros_asistencia(),
            self.lbl_asistencia.config(text="", fg=C["text_secondary"])))

    def _calcular_estado(self, entrada, salida):
        """Clasifica la marcación del día según entrada, salida y horarios.

        Args:
            entrada (str | None): Hora de entrada (HH:MM) o None.
            salida (str | None): Hora de salida (HH:MM) o None.

        Returns:
            str: Inasistencia, Atrasado, Salida anticipada, Pendiente de salida,
                 Solo salida, Atrasado y salida anticipada o Normal.
        """
        if not entrada and not salida:
            return "Inasistencia"
        if not entrada:
            return "Solo salida"
        if entrada > self.hora_entrada:
            if salida and salida < self.hora_salida:
                return "Atrasado y salida anticipada"
            return "Atrasado"
        if salida and salida < self.hora_salida:
            return "Salida anticipada"
        if not salida:
            return "Pendiente de salida"
        return "Normal"

    # ------------------------------------------------------------------
    # REPORTES
    # ------------------------------------------------------------------
    def _crear_reportes(self):
        """Construye la página Reportes: filtros, métricas, pestañas y exportación."""
        frame = tk.Frame(self._scroll_frame, bg=C["bg_app"])
        self.paginas["reportes"] = frame

        PageHeader(frame, "Reportes",
                   "Atrasos, salidas anticipadas, inasistencias y alertas."
                   ).pack(fill="x", padx=24, pady=(20, 10))

        barra = tk.Frame(frame, bg=C["bg_app"])
        barra.pack(fill="x", padx=24, pady=(0, 10))

        self.entry_desde = self._combo_fecha(barra, (datetime.now() -
                                                    timedelta(days=6)).strftime("%Y-%m-%d"))
        self.entry_desde.pack(side="left", padx=(0, 10))
        self.entry_hasta = self._combo_fecha(barra, self._obtener_fecha_hoy())
        self.entry_hasta.pack(side="left", padx=(0, 10))
        for campo in (self.entry_desde, self.entry_hasta):
            campo.bind("<<ComboboxSelected>>",
                       lambda e: self._programar_generar_reportes())

        self.entry_buscar_reporte = tk.Entry(
            barra, bg=C["input_bg"], fg=C["text_secondary"], relief="flat",
            insertbackground=C["text_primary"], font=F["body"],
            highlightthickness=1, highlightbackground=C["border"],
            highlightcolor=C["border_active"], width=22)
        self.entry_buscar_reporte.insert(0, "Buscar trabajador...")
        self.entry_buscar_reporte.pack(side="left", padx=(0, 10), ipady=9)
        self.entry_buscar_reporte.bind("<FocusIn>", self._limpiar_placeholder_reporte)
        self.entry_buscar_reporte.bind(
            "<KeyRelease>", lambda e: self._aplicar_filtro_reporte())

        self.btn_generar_reportes = PrimaryButton(barra, "GENERAR REPORTES", icono="grafica")
        self.btn_generar_reportes.pack(side="left", padx=(0, 10))
        self.btn_generar_reportes.config(command=self._lanzar_generar_reportes)

        self.btn_exportar = SecondaryButton(barra, "EXPORTAR EXCEL", icono=None)
        self.btn_exportar.pack(side="left")
        self.btn_exportar.config(command=self._exportar_excel)

        fila_estado = tk.Frame(frame, bg=C["bg_app"])
        fila_estado.pack(fill="x", padx=24, pady=(0, 10))
        self.lbl_reporte = tk.Label(fila_estado, text="", bg=C["bg_app"],
                                    fg=C["text_secondary"], font=F["small"],
                                    anchor="w", justify="left", wraplength=1000)
        self.lbl_reporte.pack(fill="x")

        fila_metricas = tk.Frame(frame, bg=C["bg_app"])
        fila_metricas.pack(fill="x", padx=24, pady=(0, 10))

        self.metricas_reporte = {}
        defs = [
            ("total", "Total incidencias"),
            ("pendientes", "Pendientes"),
            ("resueltas", "Resueltas"),
        ]
        for i, (clave, texto) in enumerate(defs):
            card = tk.Frame(fila_metricas, bg=C["surface"], highlightthickness=1,
                            highlightbackground=C["border"], padx=16, pady=10)
            card.pack(side="left", expand=True, fill="x",
                      padx=(0 if i == 0 else 8, 0 if i == 2 else 8))
            tk.Label(card, text=texto.upper(), bg=C["surface"],
                     fg=C["text_secondary"], font=F["small_bold"]).pack(anchor="w")
            self.metricas_reporte[clave] = tk.Label(
                card, text="0", bg=C["surface"], fg=C["text_primary"],
                font=F["metric_value"], anchor="w")
            self.metricas_reporte[clave].pack(anchor="w")

        notebook = ttk.Notebook(frame, style="Dark.TNotebook")
        notebook.pack(fill="both", expand=True, padx=24, pady=(0, 14))
        self.notebook_reportes = notebook

        def _tab(titulo, clave, cols_def, con_check=False):
            """Crea una pestaña del notebook de reportes con su tabla y scroll.

            Args:
                titulo (str): Texto de la pestaña.
                clave (str): Clave del tipo de reporte.
                cols_def (list): Definiciones (columna, título, ancho).
                con_check (bool): Si incluye columna de checkboxes.

            Returns:
                ttk.Treeview: La tabla creada.
            """
            tab = tk.Frame(notebook, bg=C["surface"])
            cont = tk.Frame(tab, bg=C["surface"])
            cont.pack(fill="both", expand=True, padx=6, pady=6)
            show = "tree headings" if con_check else "headings"
            tabla = ttk.Treeview(
                cont, columns=[c[0] for c in cols_def], show=show,
                selectmode="extended" if con_check else "browse",
                style="Dark.Treeview")
            if con_check:
                tabla.heading("#0", text="")
                tabla.column("#0", width=40, anchor="center")
                tabla.tag_configure("checked", foreground=C["text_primary"])
                tabla.bind("<Button-1>", self._toggle_check_alerta)
            for i, (c, tit, ancho) in enumerate(cols_def):
                tabla.heading(c, text=tit,
                              command=lambda cc=c, ci=i, k=clave: self._ordenar_reporte(k, cc, ci))
                tabla.column(c, width=ancho)
            tabla.pack(side="left", fill="both", expand=True)
            scroll = ttk.Scrollbar(cont, orient="vertical", command=tabla.yview,
                                   style="Dark.Scrollbar")
            scroll.pack(side="right", fill="y")
            tabla.configure(yscrollcommand=scroll.set)
            notebook.add(tab, text=titulo)
            return tabla

        self.tabla_atrasos = _tab("Atrasos", "atrasos", [
            ("trabajador", "Trabajador", 240),
            ("fecha", "Fecha", 110),
            ("hora", "Hora de entrada", 140),
            ("min", "Minutos de atraso", 150),
        ])
        self.tabla_salidas = _tab("Salidas anticipadas", "salidas", [
            ("trabajador", "Trabajador", 240),
            ("fecha", "Fecha", 110),
            ("hora", "Hora de salida", 140),
            ("min", "Minutos anticipados", 160),
        ])
        self.tabla_inasistencias = _tab("Inasistencias", "inasistencias", [
            ("trabajador", "Trabajador", 240),
            ("fecha", "Fecha", 110),
            ("entrada", "Entrada", 110),
            ("salida", "Salida", 110),
            ("estado", "Estado", 130),
        ])
        self.tabla_alertas = _tab("Alertas", "alertas", [
            ("trabajador", "Trabajador", 220),
            ("tipo", "Tipo", 180),
            ("fecha", "Fecha", 110),
            ("hora", "Hora", 100),
            ("estado", "Estado", 110),
        ], con_check=True)

        self._alertas_por_fila = {}
        self._checkeos_reportes = set()
        self._cache_reportes = {
            "atrasos": [], "salidas": [], "inasistencias": [], "alertas": [],
        }
        self._populadores = {
            "atrasos": self.tabla_atrasos,
            "salidas": self.tabla_salidas,
            "inasistencias": self.tabla_inasistencias,
            "alertas": self.tabla_alertas,
        }
        self.orden_reporte = {}

        barra_acciones = tk.Frame(frame, bg=C["bg_app"])
        barra_acciones.pack(fill="x", padx=24, pady=(0, 16))
        SecondaryButton(barra_acciones, "MARCAR COMO LE\u00cdDA",
                        command=self._marcar_alerta_leida).pack(side="left")

    def _combo_fecha(self, padre, valor):
        """Crea un desplegable de fechas con valores solo hasta hoy.

        Incluye los últimos 60 días y excluye fechas futuras para que el
        rango de reportes no apunte a días sin marcaciones.
        """
        hoy = datetime.now()
        opciones = [(hoy - timedelta(days=i)).strftime("%Y-%m-%d")
                    for i in range(60, -1, -1)]
        combo = ttk.Combobox(
            padre, values=opciones, state="readonly", style="Dark.TCombobox",
            font=F["body"], width=12, justify="center")
        combo.set(valor)
        return combo

    def _limpiar_placeholder_reporte(self, e):
        """Borra el texto de ejemplo del buscador de reportes al recibir el foco."""
        if self.entry_buscar_reporte.get() == "Buscar trabajador...":
            self.entry_buscar_reporte.delete(0, "end")
        self.entry_buscar_reporte.config(fg=C["text_primary"])

    def _toggle_check_alerta(self, event):
        """Marca o desmarca una alerta de la tabla cuando se hace clic en su fila."""
        region = self.tabla_alertas.identify("region", event.x, event.y)
        if region != "cell":
            return
        fila_id = self.tabla_alertas.identify_row(event.y)
        if not fila_id:
            return
        valores = tuple(self.tabla_alertas.item(fila_id, "values"))
        if valores in self._checkeos_reportes:
            self._checkeos_reportes.discard(valores)
            self.tabla_alertas.item(fila_id, text="\u2610")
            self.tabla_alertas.tag_configure("checked", foreground=C["text_primary"])
        else:
            self._checkeos_reportes.add(valores)
            self.tabla_alertas.item(fila_id, text="\u2612")
        return "break"

    def _aplicar_filtro_reporte(self):
        """Filtra los reportes cacheados por el texto del buscador y actualiza la UI."""
        texto = self.entry_buscar_reporte.get().strip().lower()
        if texto == "buscar trabajador...":
            texto = ""
        for clave, cache in self._cache_reportes.items():
            filas = cache
            if texto:
                filas = [f for f in cache if texto in f[0].lower()]
            self._poblar_tabla(self._populadores[clave], filas)
        self._actualizar_titulos_tabs()
        self._actualizar_metricas()

    def _poblar_tabla(self, tabla, filas):
        """Rellena una tabla de reportes y muestra un aviso cuando no hay datos."""
        for item in tabla.get_children():
            tabla.delete(item)
        if not filas:
            tabla.insert("", "end", values=("Sin datos en el rango seleccionado",))
            return
        for f in filas:
            tabla.insert("", "end", values=f, text="\u2610")

    def _ordenar_reporte(self, clave, col_id, col_idx):
        """Alterna el orden de una pestaña por la columna indicada."""
        filas = self._cache_reportes.get(clave)
        if not filas:
            return
        self.orden_reporte[clave] = not self.orden_reporte.get(clave, False)
        ordenadas = sorted(filas, key=lambda f: str(f[col_idx]).lower(),
                           reverse=self.orden_reporte[clave])
        self.ventana.after(0, lambda: self._poblar_tabla(self._populadores[clave],
                                                         ordenadas))

    def _actualizar_titulos_tabs(self):
        """Actualiza los títulos de las pestañas con el conteo de filas visibles."""
        nombres = {
            "atrasos": "Atrasos",
            "salidas": "Salidas anticipadas",
            "inasistencias": "Inasistencias",
            "alertas": "Alertas",
        }
        texto = self._filtro_activo()
        for idx, (clave, titulo) in enumerate(nombres.items()):
            filas = self._cache_reportes[clave]
            if texto:
                filas = [f for f in filas if texto in f[0].lower()]
            self.notebook_reportes.tab(idx, text=f"{titulo} ({len(filas)})")

    def _actualizar_metricas(self):
        """Recalcula las tarjetas de métricas a partir de las alertas cacheadas."""
        alertas = self._cache_reportes.get("alertas", [])
        total = len(alertas)
        pendientes = len([a for a in alertas if "Pendiente" in a[4]])
        resueltas = total - pendientes
        self.metricas_reporte["total"].config(text=str(total))
        self.metricas_reporte["pendientes"].config(text=str(pendientes))
        self.metricas_reporte["resueltas"].config(text=str(resueltas))

    def _filtro_activo(self):
        """Texto de búsqueda vigente en el buscador de reportes (o cadena vacía)."""
        texto = self.entry_buscar_reporte.get().strip().lower()
        if texto == "buscar trabajador...":
            return ""
        return texto

    def _programar_generar_reportes(self):
        """Agenda la regeneración de reportes tras cambiar el rango de fechas."""
        if getattr(self, "_after_reportes_id", None):
            self.ventana.after_cancel(self._after_reportes_id)
        self._after_reportes_id = self.ventana.after(
            650, self._lanzar_generar_reportes)

    def _lanzar_generar_reportes(self):
        """Lanza la generación de reportes en un hilo (evita recargas dobles)."""
        if getattr(self, "_after_reportes_id", None):
            self.ventana.after_cancel(self._after_reportes_id)
            self._after_reportes_id = None
        if not self.firebase_listo or getattr(self, "_generando_reportes", False):
            return
        self._generando_reportes = True
        threading.Thread(target=self._hilo_generar_reportes, daemon=True).start()

    def _hilo_generar_reportes(self):
        """Ejecuta la generación de reportes y libera el candado al terminar."""
        try:
            self._generar_reportes()
        finally:
            self._generando_reportes = False

    def _generar_reportes(self):
        """Valida el rango y genera los cuatro reportes (atrasos, salidas, etc.)."""
        desde = self.entry_desde.get().strip()
        hasta = self.entry_hasta.get().strip()
        try:
            d = datetime.strptime(desde, "%Y-%m-%d")
            h = datetime.strptime(hasta, "%Y-%m-%d")
        except Exception:
            self.ventana.after(0, lambda: self.lbl_reporte.config(
                text="Rango de fechas inv\u00e1lido (usa AAAA-MM-DD).", fg=C["danger"]))
            return
        if d > h:
            d, h = h, d
        if h > datetime.now():
            self.ventana.after(0, lambda: self.lbl_reporte.config(
                text="El rango no puede incluir fechas futuras.", fg=C["danger"]))
            return
        desde = d.strftime("%Y-%m-%d")
        hasta = h.strftime("%Y-%m-%d")
        dias = set()
        dia = d
        while dia <= h:
            dias.add(dia.strftime("%Y-%m-%d"))
            dia += timedelta(days=1)

        self.ventana.after(0, lambda: (
            self.btn_generar_reportes.set_loading(True, "GENERANDO..."),
            self.lbl_reporte.config(text="Generando...", fg=C["text_secondary"])))
        self._generar_atrasos(desde, hasta)
        self._generar_salidas(desde, hasta)
        self._generar_inasistencias(dias)
        self._generar_alertas(desde, hasta)
        self.ventana.after(0, lambda: (
            self.btn_generar_reportes.set_loading(False),
            self.lbl_reporte.config(text="Reportes generados.", fg=C["success"]),
            self._aplicar_filtro_reporte()))

    def _generar_atrasos(self, desde, hasta):
        """Calcula los atrasos (entradas posteriores a la hora límite) del rango."""
        try:
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", ">=", desde)) \
                .where(filter=FieldFilter("fecha", "<=", hasta)).get()
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_reporte.config(
                text=str(e), fg=C["danger"]))
            return
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in self.usuarios}
        filas = []
        for d in docs:
            m = d.to_dict()
            if m.get("tipo") != "entrada":
                continue
            hora = (m.get("hora") or "")[:5]
            if hora and hora > self.hora_entrada:
                filas.append((
                    nombres.get(m.get("usuario"), m.get("usuario") or "?"),
                    m.get("fecha", ""),
                    hora,
                    _minutos_diff(self.hora_entrada, hora),
                ))
        filas.sort(key=lambda r: (r[0], r[1]))
        self._cache_reportes["atrasos"] = filas
        self.ventana.after(0, lambda: self._poblar_tabla(self.tabla_atrasos, filas))

    def _generar_salidas(self, desde, hasta):
        """Calcula las salidas anticipadas (anteriores a la hora límite) del rango."""
        try:
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", ">=", desde)) \
                .where(filter=FieldFilter("fecha", "<=", hasta)).get()
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_reporte.config(
                text=str(e), fg=C["danger"]))
            return
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in self.usuarios}
        filas = []
        for d in docs:
            m = d.to_dict()
            if m.get("tipo") != "salida":
                continue
            hora = (m.get("hora") or "")[:5]
            if hora and hora < self.hora_salida:
                filas.append((
                    nombres.get(m.get("usuario"), m.get("usuario") or "?"),
                    m.get("fecha", ""),
                    hora,
                    _minutos_diff(hora, self.hora_salida),
                ))
        filas.sort(key=lambda r: (r[0], r[1]))
        self._cache_reportes["salidas"] = filas
        self.ventana.after(0, lambda: self._poblar_tabla(self.tabla_salidas, filas))

    def _generar_inasistencias(self, dias):
        """Lista los días hábiles sin registro de cada trabajador en el rango."""
        try:
            desde, hasta = min(dias), max(dias)
            docs = self.db.collection("marcaciones") \
                .where(filter=FieldFilter("fecha", ">=", desde)) \
                .where(filter=FieldFilter("fecha", "<=", hasta)).get()
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_reporte.config(
                text=str(e), fg=C["danger"]))
            return
        registrados = {}
        for d in docs:
            m = d.to_dict()
            registrados.setdefault((m.get("usuario"), m.get("fecha")), True)
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in self.usuarios}
        filas = []
        hoy = self._obtener_fecha_hoy()
        for u in self._trabajadores():
            for fecha in dias:
                if not _es_dia_habl(fecha) or fecha > hoy:
                    continue
                if (u.get("usuario"), fecha) not in registrados:
                    filas.append((
                        nombres.get(u.get("usuario"), u.get("usuario")),
                        fecha,
                        "Sin registro",
                        "Sin registro",
                        "Inasistencia",
                    ))
        filas.sort(key=lambda r: (r[0], r[1]))
        self._cache_reportes["inasistencias"] = filas
        self.ventana.after(0, lambda: self._poblar_tabla(self.tabla_inasistencias,
                                                         filas))

    def _generar_alertas(self, desde, hasta):
        """Carga las alertas del rango y las asocia con su id para marcarlas leídas."""
        try:
            alertas = Alerta.listar(self.db, fecha_inicio=desde, fecha_fin=hasta)
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_reporte.config(
                text=str(e), fg=C["danger"]))
            return
        nombres = {u.get("usuario"): (u.get("nombre") or u.get("usuario"))
                   for u in self.usuarios}
        etiquetas = {
            Alerta.TIPO_ATRASO: "Atraso",
            Alerta.TIPO_SALIDA_ANTICIPADA: "Salida anticipada",
            Alerta.TIPO_INASISTENCIA: "Inasistencia",
        }
        filas = []
        mapa = {}
        for a in alertas:
            fila = (
                nombres.get(a.get("usuario"), a.get("usuario") or "?"),
                etiquetas.get(a.get("tipo"), a.get("tipo")),
                a.get("fecha", ""),
                a.get("hora") or "",
                "Le\u00edda" if a.get("estado") == Alerta.ESTADO_LEIDA else "Pendiente",
            )
            filas.append(fila)
            mapa[fila] = a.get("_id")
        self._alertas_por_fila = mapa
        self._cache_reportes["alertas"] = filas
        self.ventana.after(0, lambda: self._poblar_tabla(self.tabla_alertas, filas))

    def _marcar_alerta_leida(self):
        """Marca como leídas las alertas seleccionadas con los checkboxes."""
        a_marcar = []
        for fila in list(self._checkeos_reportes):
            alerta_id = self._alertas_por_fila.get(tuple(fila))
            if alerta_id:
                a_marcar.append(alerta_id)
        if not a_marcar:
            self._mostrar_toast("Selecciona una alerta para marcarla como le\u00edda",
                                ok=False)
            return
        self._checkeos_reportes.clear()
        threading.Thread(target=self._marcar_alertas_leidas,
                         args=(a_marcar,), daemon=True).start()

    def _exportar_excel(self):
        """Exporta todos los reportes a un archivo .xlsx con formato."""
        claves = ("atrasos", "salidas", "inasistencias", "alertas")
        if not any(self._cache_reportes.get(k) for k in claves):
            self._mostrar_toast(
                "No hay datos para exportar. Genera los reportes primero.",
                ok=False)
            return

        nombre = filedialog.asksaveasfilename(
            title="Exportar a Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"reporte_asistencia_{self._obtener_fecha_hoy()}.xlsx")
        if not nombre:
            return

        def _construir_y_guardar():
            """Construye el libro de Excel y lo guarda sin bloquear la interfaz."""
            from openpyxl import Workbook

            encabezados = {
                "atrasos": ["Trabajador", "Fecha", "Hora de entrada",
                            "Minutos de atraso"],
                "salidas": ["Trabajador", "Fecha", "Hora de salida",
                            "Minutos anticipados"],
                "inasistencias": ["Trabajador", "Fecha", "Entrada", "Salida",
                                  "Estado"],
                "alertas": ["Trabajador", "Tipo", "Fecha", "Hora", "Estado"],
            }
            titulos = {
                "atrasos": "Atrasos",
                "salidas": "Salidas anticipadas",
                "inasistencias": "Inasistencias",
                "alertas": "Alertas",
            }
            estilos_estado = {
                "Inasistencia": C["danger"].lstrip("#"),
                "Atrasado": C["warning"].lstrip("#"),
                "Salida anticipada": C["warning"].lstrip("#"),
                "Pendiente": C["warning"].lstrip("#"),
                "Le\u00edda": C["success"].lstrip("#"),
                "Normal": C["success"].lstrip("#"),
            }
            try:
                wb = Workbook()
                wb.remove(wb.active)
                self._hoja_resumen(wb, titulos)
                for clave in claves:
                    self._hoja_reporte(wb, clave, encabezados[clave],
                                       titulos[clave], estilos_estado)
                wb.save(nombre)
                self.ventana.after(
                    0, lambda: self._mostrar_toast("Excel exportado correctamente"))
            except Exception as e:
                self.ventana.after(0, lambda e=e: self._mostrar_toast(
                    f"Error al exportar: {e}", ok=False))

        threading.Thread(target=_construir_y_guardar, daemon=True).start()

    def _hoja_resumen(self, wb, titulos):
        """Agrega la hoja Resumen con los totales por tipo de reporte."""
        from openpyxl.styles import Alignment, Font, PatternFill

        hoja = wb.create_sheet("Resumen")
        hoja.append(["Reporte", "Cantidad"])
        filas = [
            (titulos["atrasos"], len(self._cache_reportes["atrasos"])),
            (titulos["salidas"], len(self._cache_reportes["salidas"])),
            (titulos["inasistencias"], len(self._cache_reportes["inasistencias"])),
            ("Alertas", len(self._cache_reportes["alertas"])),
        ]
        n_alertas = len(self._cache_reportes["alertas"])
        pendientes = sum(1 for a in self._cache_reportes["alertas"]
                         if "Pendiente" in a[4])
        filas.append(("Alertas pendientes", pendientes))
        filas.append(("Alertas resueltas", n_alertas - pendientes))
        for fila in filas:
            hoja.append(list(fila))

        color = C["accent"].lstrip("#")
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=color)
            celda.alignment = Alignment(horizontal="center")
        hoja.freeze_panes = "A2"
        hoja.column_dimensions["A"].width = 24
        hoja.column_dimensions["B"].width = 14

    def _hoja_reporte(self, wb, clave, encabezado, titulo, estilos_estado):
        """Agrega una hoja de reporte con encabezado, anchos, filtro y colores."""
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        hoja = wb.create_sheet(titulo[:31])
        hoja.append(encabezado)
        for fila in self._cache_reportes.get(clave, []):
            hoja.append(list(fila))

        n_columnas = len(encabezado)
        color = C["accent"].lstrip("#")
        for i, celda in enumerate(hoja[1], 1):
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=color)
            celda.alignment = Alignment(horizontal="center")
            letra = get_column_letter(i)
            largos = [len(str(celda.value))]
            for fila in hoja.iter_rows(min_col=i, max_col=i, min_row=2):
                for celdilla in fila:
                    if celdilla.value is not None:
                        if isinstance(celdilla.value, int):
                            celdilla.alignment = Alignment(horizontal="center")
                        largos.append(len(str(celdilla.value)))
            hoja.column_dimensions[letra].width = min(max(largos) + 2, 42)

        estado_col = next((i for i, h in enumerate(encabezado)
                           if h == "Estado"), None)
        if estado_col is not None:
            for fila in hoja.iter_rows(min_row=2, max_row=hoja.max_row):
                celda = fila[estado_col]
                texto = str(celda.value or "")
                for palabra, hex_color in estilos_estado.items():
                    if palabra in texto:
                        celda.font = Font(bold=True, color=hex_color)
                        break

        if n_columnas:
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions

    # ------------------------------------------------------------------
    # CONFIGURACI\u00d3N
    # ------------------------------------------------------------------
    def _crear_config(self):
        """Construye la página Configuración: empresa, horarios y guardado."""
        frame = tk.Frame(self._scroll_frame, bg=C["bg_app"])
        self.paginas["config"] = frame

        PageHeader(frame, "Configuraci\u00f3n",
                   "Par\u00e1metros generales del sistema."
                   ).pack(fill="x", padx=24, pady=(20, 12))

        fila = tk.Frame(frame, bg=C["bg_app"])
        fila.pack(fill="x", padx=24)
        fila.columnconfigure(0, weight=1)
        fila.columnconfigure(1, weight=1)

        card_empresa = SurfaceCard(fila, padding=20)
        card_empresa.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        tk.Label(card_empresa, text="DATOS DE LA EMPRESA", bg=C["surface"],
                 fg=C["text_secondary"], font=F["small_bold"]).pack(anchor="w")
        tk.Label(card_empresa, text="Nombre visible en los reportes.",
                 bg=C["surface"], fg=C["text_muted"], font=F["small"]).pack(
            anchor="w", pady=(0, 10))

        tk.Label(card_empresa, text="Nombre de la empresa", bg=C["surface"],
                 fg=C["text_primary"], font=F["body_bold"]).pack(anchor="w", pady=(4, 0))
        self.entry_empresa = tk.Entry(
            card_empresa, bg=C["input_bg"], fg=C["text_primary"], relief="flat",
            insertbackground=C["text_primary"], font=F["body"],
            highlightthickness=1, highlightbackground=C["border"],
            highlightcolor=C["border_active"])
        self.entry_empresa.pack(fill="x", pady=(4, 2), ipady=9)

        self.lbl_err_empresa = tk.Label(card_empresa, text="", bg=C["surface"],
                                        fg=C["danger"], font=F["small"])
        self.lbl_err_empresa.pack(anchor="w")

        tk.Label(card_empresa, text="Notificaciones", bg=C["surface"],
                 fg=C["text_primary"], font=F["body_bold"]).pack(
            anchor="w", pady=(14, 0))
        tk.Label(card_empresa, text="Recibir avisos de inasistencia.",
                 bg=C["surface"], fg=C["text_muted"], font=F["small"]).pack(
            anchor="w", pady=(2, 4))
        self.var_notif = tk.BooleanVar(value=False)
        tk.Checkbutton(card_empresa, text="Activar notificaciones",
                       variable=self.var_notif, bg=C["surface"], fg=C["text_primary"],
                       selectcolor=C["surface_raised"], activebackground=C["surface"],
                       activeforeground=C["text_primary"], font=F["body"],
                       highlightthickness=0, bd=0).pack(anchor="w")

        card_horarios = SurfaceCard(fila, padding=20)
        card_horarios.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        tk.Label(card_horarios, text="HORARIOS", bg=C["surface"],
                 fg=C["text_secondary"], font=F["small_bold"]).pack(anchor="w")
        tk.Label(card_horarios, text="L\u00edmites para considerar atraso y salida anticipada.",
                 bg=C["surface"], fg=C["text_muted"], font=F["small"]).pack(
            anchor="w", pady=(0, 10))

        grid = tk.Frame(card_horarios, bg=C["surface"])
        grid.pack(fill="x")
        grid.columnconfigure(0, weight=1)

        tk.Label(grid, text="Hora l\u00edmite de entrada",
                 bg=C["surface"], fg=C["text_primary"], font=F["body_bold"]).grid(
            row=0, column=0, sticky="w", pady=(4, 0))
        self.entry_hora_ent = self._entry_hora(grid, "09:30")
        self.entry_hora_ent.grid(row=1, column=0, sticky="ew", pady=(4, 2), ipady=9)
        self.lbl_err_ent = tk.Label(grid, text="", bg=C["surface"], fg=C["danger"],
                                    font=F["small"])
        self.lbl_err_ent.grid(row=2, column=0, sticky="w")

        tk.Label(grid, text="Hora l\u00edmite de salida",
                 bg=C["surface"], fg=C["text_primary"], font=F["body_bold"]).grid(
            row=3, column=0, sticky="w", pady=(8, 0))
        self.entry_hora_sal = self._entry_hora(grid, "17:30")
        self.entry_hora_sal.grid(row=4, column=0, sticky="ew", pady=(4, 2), ipady=9)
        self.lbl_err_sal = tk.Label(grid, text="", bg=C["surface"], fg=C["danger"],
                                    font=F["small"])
        self.lbl_err_sal.grid(row=5, column=0, sticky="w")

        barra_inferior = tk.Frame(frame, bg=C["bg_app"])
        barra_inferior.pack(fill="x", padx=24, pady=(14, 20))
        self.lbl_ultimo_guardado = tk.Label(
            barra_inferior, text="A\u00fan no hay cambios guardados",
            bg=C["bg_app"], fg=C["text_muted"], font=F["small"])
        self.lbl_ultimo_guardado.pack(side="left")
        SecondaryButton(barra_inferior, "DESCARTAR CAMBIOS",
                        command=self._descartar_config).pack(side="right")
        self.btn_guardar_config = PrimaryButton(barra_inferior, "GUARDAR CAMBIOS")
        self.btn_guardar_config.pack(side="right", padx=(0, 8))
        self.btn_guardar_config.config(
            command=lambda: threading.Thread(target=self._guardar_config, daemon=True).start())

    def _entry_hora(self, padre, valor):
        """Crea un campo de entrada para horas con un valor inicial."""
        entry = tk.Entry(
            padre, bg=C["input_bg"], fg=C["text_primary"], relief="flat",
            insertbackground=C["text_primary"], font=F["body"], width=10,
            highlightthickness=1, highlightbackground=C["border"],
            highlightcolor=C["border_active"], justify="center")
        entry.insert(0, valor)
        return entry

    def _poblar_config_ui(self):
        """Vuelca los valores actuales de configuración en los campos del formulario."""
        if getattr(self, "_cerrando", False) or not hasattr(self, "entry_empresa"):
            return
        self.entry_empresa.delete(0, "end")
        self.entry_empresa.insert(0, self.nombre_empresa)
        self.entry_hora_ent.delete(0, "end")
        self.entry_hora_ent.insert(0, self.hora_entrada)
        self.entry_hora_sal.delete(0, "end")
        self.entry_hora_sal.insert(0, self.hora_salida)
        self.var_notif.set(self.notificaciones)

    def _descartar_config(self):
        """Restaura los campos de configuración a sus valores guardados."""
        self._poblar_config_ui()
        for lbl in (self.lbl_err_empresa, self.lbl_err_ent, self.lbl_err_sal):
            lbl.config(text="")
        self.lbl_ultimo_guardado.config(
            text="Cambios descartados", fg=C["text_muted"])
        self.ventana.after(2500, lambda: self.lbl_ultimo_guardado.config(
            text="A\u00fan no hay cambios guardados"))

    def _guardar_config(self):
        """Valida y guarda la configuración de empresa y horarios en Firestore."""
        nombre = self.entry_empresa.get().strip()
        hora_ent = self.entry_hora_ent.get().strip()
        hora_sal = self.entry_hora_sal.get().strip()

        def _error(lbl, msg):
            """Muestra un error en la etiqueta dada y quita el estado de carga."""
            self.ventana.after(0, lambda: (lbl.config(text=msg),
                                           self.btn_guardar_config.set_loading(False)))

        if not nombre:
            _error(self.lbl_err_empresa, "El nombre de la empresa es obligatorio.")
            return
        if not _hora_valida(hora_ent):
            _error(self.lbl_err_ent, "Formato HH:MM (ej: 09:30).")
            return
        if not _hora_valida(hora_sal):
            _error(self.lbl_err_sal, "Formato HH:MM (ej: 17:30).")
            return
        if hora_ent >= hora_sal:
            _error(self.lbl_err_sal, "Debe ser posterior a la hora de entrada.")
            return

        for lbl in (self.lbl_err_empresa, self.lbl_err_ent, self.lbl_err_sal):
            lbl.config(text="")

        datos = {
            "nombre": nombre,
            "hora_entrada": hora_ent,
            "hora_salida": hora_sal,
            "notificaciones": self.var_notif.get(),
        }
        try:
            self.db.collection("config").document("empresa").set(datos)
            self.hora_entrada = hora_ent
            self.hora_salida = hora_sal
            self.nombre_empresa = nombre
            self.notificaciones = self.var_notif.get()
            ahora = datetime.now().strftime("%d/%m/%Y %H:%M")
            self.ventana.after(0, lambda: (
                self.lbl_ultimo_guardado.config(
                    text=f"\u00daltima actualizaci\u00f3n: {ahora}",
                    fg=C["success"]),
                self.btn_guardar_config.set_loading(False),
                self._mostrar_toast("Configuraci\u00f3n guardada")))
        except Exception as e:
            self.ventana.after(0, lambda e=e: self.lbl_ultimo_guardado.config(
                text=f"Error al guardar: {e}", fg=C["danger"]))

    # ------------------------------------------------------------------
    # USUARIOS
    # ------------------------------------------------------------------
    def _abrir_gestion_usuarios(self):
        """Abre la ventana de gestión de usuarios."""
        import gestion_usuarios
        gestion_usuarios.GestionUsuariosApp(self.ventana, self.usuario)

    # ------------------------------------------------------------------
    # CERRAR SESI\u00d3N / ANIMACIONES
    # ------------------------------------------------------------------
    def _cerrar_sesion(self):
        """Cierra la ventana y vuelve a abrir la ventana de login."""
        self._cerrando = True
        self._detener_polling()
        self.volver_login = True
        self.ventana.destroy()

    def _mostrar_toast(self, texto, ok=True):
        """Muestra una notificación breve de éxito o error."""
        toast = Toast(self.ventana, texto, tipo="exito" if ok else "error")
        toast.mostrar()

    def _iniciar_arrastre(self, e):
        """Guarda el punto inicial para arrastrar la ventana sin bordes."""
        self._offset_x = e.x
        self._offset_y = e.y

    def _arrastrar(self, e):
        """Mueve la ventana según el desplazamiento del puntero."""
        x = self.ventana.winfo_x() + e.x - self._offset_x
        y = self.ventana.winfo_y() + e.y - self._offset_y
        self.ventana.geometry(f"+{x}+{y}")


if __name__ == "__main__":
    DashboardAdminApp("admin", {
        "usuario": "admin",
        "clave": "admin123",
        "rol": "administrador",
        "nombre": "Administrador",
        "correo": "admin@empresa.com",
    })